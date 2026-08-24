"""OOXML structure assertions used before a generated document is published."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from fire_inspect.templates import validate_docx

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


class DocumentStructureError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class DocumentStructureReport:
    document_type: str
    package_part_count: int
    checks: tuple[str, ...]


_GEOMETRY = {
    "record": ("11906", "16838", "595", "822", "369", "822"),
    "photo": ("12240", "15840", "1440", "1800", "1440", "1800"),
    "interview": ("11900", "16840", "1032", "1117", "1228", "1120"),
}


def _geometry(root: etree._Element) -> tuple[str | None, ...]:
    size = root.xpath(".//w:body/w:sectPr/w:pgSz", namespaces=NS)
    margins = root.xpath(".//w:body/w:sectPr/w:pgMar", namespaces=NS)
    if len(size) != 1 or len(margins) != 1:
        return ()
    size_element, margin_element = size[0], margins[0]
    return (
        size_element.get(qn("w:w")),
        size_element.get(qn("w:h")),
        margin_element.get(qn("w:top")),
        margin_element.get(qn("w:right")),
        margin_element.get(qn("w:bottom")),
        margin_element.get(qn("w:left")),
    )


def assert_document_structure(
    payload: bytes,
    *,
    expected_type: str,
    expected_item_count: int | None = None,
) -> DocumentStructureReport:
    validation = validate_docx(payload, filename="generated.docx")
    if expected_type not in _GEOMETRY:
        raise ValueError("未知文书类型")
    issues: list[str] = []
    checks: list[str] = []
    document = Document(BytesIO(payload))

    with ZipFile(BytesIO(payload)) as archive:
        root = etree.fromstring(
            archive.read("word/document.xml"),
            parser=etree.XMLParser(resolve_entities=False, no_network=True),
        )
        if _geometry(root) != _GEOMETRY[expected_type]:
            issues.append("纸型或页边距偏离原模板")
        else:
            checks.append("template-page-geometry")

        if expected_type == "record":
            if len(document.tables) != 1:
                issues.append("记录表必须保留一个表格")
            else:
                table = document.tables[0]
                if len(table.rows) != 5 or len(table.columns) != 5:
                    issues.append("记录表必须保留原 5×5 表格")
                checkbox_text = table.cell(0, 0).text
                if checkbox_text.count("☑") != 1 or checkbox_text.count("□") != 6:
                    issues.append("检查类型不是一个选中项和六个未选项")
                if 'w:hRule="exact"' in table._tbl.xml:
                    issues.append("记录表包含可能裁切长文本的固定行高")
            if (
                len(document.paragraphs) < 9
                or document.paragraphs[6].text
                or document.paragraphs[8].text
            ):
                issues.append("记录表签名值区域未保持空白")
            checks.extend(("record-table", "record-checkboxes", "signature-blank"))

        elif expected_type == "photo":
            count = (
                expected_item_count
                if expected_item_count is not None
                else len(document.inline_shapes)
            )
            if len(document.inline_shapes) != count:
                issues.append("照片数量与证据区块数量不一致")
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            if "[此处插入照片]" in text:
                issues.append("照片文书残留图片占位文本")
            page_breaks = root.xpath('.//w:br[@w:type="page"]', namespaces=NS)
            if len(page_breaks) != max(count - 1, 0):
                issues.append("照片证据区块之间缺少显式分页")
            for shape in document.inline_shapes:
                if int(shape.width) > int(6 * 914400) or int(shape.height) > int(5 * 914400):
                    issues.append("照片超出模板图片区域")
            checks.extend(("photo-count", "photo-placeholder-removed", "photo-page-breaks"))

        elif expected_type == "interview":
            if document.tables:
                issues.append("询问笔录正文不得改造成固定高度表格")
            qa_paragraphs = [
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith(("问：", "答："))
            ]
            if expected_item_count is not None and len(qa_paragraphs) != expected_item_count:
                issues.append("询问笔录问答段落数量不一致")
            if any(
                not paragraph.runs or any(run.font.underline is not True for run in paragraph.runs)
                for paragraph in qa_paragraphs
            ):
                issues.append("询问笔录正文未按模板规则添加下划线")
            settings = etree.fromstring(archive.read("word/settings.xml"))
            if not settings.xpath('boolean(.//w:updateFields[@w:val="true"])', namespaces=NS):
                issues.append("询问笔录未设置打开时更新页码域")
            footer_names = sorted(
                name
                for name in archive.namelist()
                if name.startswith("word/footer") and name.endswith(".xml")
            )
            if len(footer_names) != 3:
                issues.append("询问笔录没有保留首页、奇数页和偶数页页脚")
            for footer_name in footer_names:
                footer = etree.fromstring(archive.read(footer_name))
                instructions = " ".join(
                    footer.xpath(".//w:instrText/text()", namespaces=NS)
                ).upper()
                simple = " ".join(footer.xpath(".//w:fldSimple/@w:instr", namespaces=NS)).upper()
                if "PAGE" not in instructions or "NUMPAGES" not in simple:
                    issues.append("询问笔录页脚 PAGE/NUMPAGES 域不完整")
            if root.xpath(".//w:framePr", namespaces=NS):
                issues.append("询问笔录包含可能裁切正文的文本框架")
            checks.extend(("interview-qa", "interview-footer-fields", "update-fields"))

    if issues:
        raise DocumentStructureError("；".join(dict.fromkeys(issues)))
    return DocumentStructureReport(
        document_type=expected_type,
        package_part_count=len(validation.package_parts),
        checks=tuple(checks),
    )
