"""Template-derived deterministic DOCX generators for the three MVP forms."""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps, UnidentifiedImageError

from fire_inspect.templates import TemplateDocumentType, recognize_template
from fire_inspect.templates.security import DocxSource, read_docx_source, validate_docx

from .domain import format_chinese_date, format_chinese_datetime
from .inputs import (
    InterviewDocumentData,
    PhotoDocumentData,
    PhotoItem,
    RecordDocumentData,
)
from .package import (
    finalize_interview_page_count,
    normalize_package_timestamps,
    strip_external_attached_template,
)

W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
MAX_IMAGE_BYTES = 20 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
PHOTO_MAX_WIDTH = int(Inches(6.0))
PHOTO_MAX_HEIGHT = int(Inches(5.0))


class DocumentGenerationError(ValueError):
    pass


def _template_document(
    source: DocxSource, expected_type: TemplateDocumentType
) -> tuple[Document, bytes]:
    payload = read_docx_source(source)
    recognition = recognize_template(source)
    if recognition.document_type is not expected_type:
        raise DocumentGenerationError("所选模板类型与文书类型不匹配")
    try:
        document = Document(BytesIO(payload))
    except Exception as exc:
        raise DocumentGenerationError("模板无法作为 Word 文档打开") from exc
    return document, payload


def _run_properties(run: object | None) -> object | None:
    if run is None:
        return None
    element = getattr(run, "_r", None)
    rpr = element.rPr if element is not None else None
    return deepcopy(rpr) if rpr is not None else None


def _clear_paragraph(paragraph: Paragraph) -> None:
    ppr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not ppr:
            paragraph._p.remove(child)


def _add_run(
    paragraph: Paragraph,
    text: str,
    *,
    properties: object | None,
    underline: bool,
):
    run = paragraph.add_run(text)
    if properties is not None:
        current = run._r.rPr
        if current is not None:
            run._r.remove(current)
        run._r.insert(0, deepcopy(properties))
    run.font.underline = underline
    return run


def _replace_paragraph(
    paragraph: Paragraph,
    parts: tuple[tuple[str, bool, object | None], ...],
) -> None:
    _clear_paragraph(paragraph)
    for text, underline, properties in parts:
        _add_run(paragraph, text, properties=properties, underline=underline)


def _value_cell(cell, value: str, *, reference_run: object) -> None:
    paragraph = cell.paragraphs[0]
    while len(cell.paragraphs) > 1:
        cell._tc.remove(cell.paragraphs[-1]._p)
    _replace_paragraph(
        paragraph,
        ((value, False, _run_properties(reference_run)),),
    )


def _select_exactly_one_checkbox(cell, selected_label: str) -> None:
    checkbox_locations: list[tuple[Paragraph, object, int, int]] = []
    selected_location: tuple[Paragraph, int] | None = None
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if "☑" in run.text:
                run.text = run.text.replace("☑", "□")
        full_text = paragraph.text
        label_index = full_text.find(selected_label)
        if label_index >= 0:
            selected_location = (paragraph, label_index)
        cursor = 0
        for run in paragraph.runs:
            for index, character in enumerate(run.text):
                if character == "□":
                    checkbox_locations.append((paragraph, run, index, cursor + index))
            cursor += len(run.text)
    if selected_location is None:
        raise DocumentGenerationError("模板中没有找到所选检查类型")
    selected_paragraph, label_index = selected_location
    candidates = [
        location
        for location in checkbox_locations
        if location[0] is selected_paragraph and location[3] <= label_index
    ]
    if not candidates:
        raise DocumentGenerationError("模板检查类型复选框结构不匹配")
    _, run, character_index, _ = max(candidates, key=lambda item: item[3])
    characters = list(run.text)
    characters[character_index] = "☑"
    run.text = "".join(characters)
    all_text = "".join(paragraph.text for paragraph in cell.paragraphs)
    if all_text.count("☑") != 1 or all_text.count("□") != 6:
        raise DocumentGenerationError("模板检查类型必须恰好包含七个选项")


def _save_document(document: Document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    payload = normalize_package_timestamps(
        strip_external_attached_template(stream.getvalue())
    )
    validate_docx(payload, filename="generated.docx")
    return payload


def generate_record_docx(source: DocxSource, data: RecordDocumentData) -> bytes:
    document, _ = _template_document(source, TemplateDocumentType.RECORD)
    if len(document.paragraphs) != 10 or len(document.tables) != 1:
        raise DocumentGenerationError("监督检查记录表模板结构与已确认映射不一致")
    table = document.tables[0]
    if len(table.rows) != 5 or len(table.columns) != 5:
        raise DocumentGenerationError("监督检查记录表的 5×5 表格结构已变化")

    number_paragraph = document.paragraphs[3]
    number_label_style = _run_properties(number_paragraph.runs[0])
    number_value_style = _run_properties(number_paragraph.runs[-2])
    _replace_paragraph(
        number_paragraph,
        (
            ("编号: ", False, number_label_style),
            (data.document_number.strip(), False, number_value_style),
        ),
    )

    date_paragraph = document.paragraphs[4]
    date_label_style = _run_properties(date_paragraph.runs[0])
    date_value_style = _run_properties(date_paragraph.runs[2])
    _replace_paragraph(
        date_paragraph,
        (
            ("检查时间: ", False, date_label_style),
            (format_chinese_date(data.inspection_date), False, date_value_style),
        ),
    )

    _select_exactly_one_checkbox(table.cell(0, 0), data.inspection_type.value)
    _value_cell(
        table.cell(1, 2),
        data.unit_name.strip(),
        reference_run=table.cell(1, 1).paragraphs[0].runs[0],
    )
    _value_cell(
        table.cell(1, 4),
        data.responsible_person.strip(),
        reference_run=table.cell(1, 3).paragraphs[0].runs[0],
    )
    _value_cell(
        table.cell(2, 2),
        data.address.strip(),
        reference_run=table.cell(2, 1).paragraphs[0].runs[0],
    )
    contact = " ".join(
        part for part in (data.contact_name.strip(), data.contact_phone.strip()) if part
    )
    _value_cell(
        table.cell(2, 4),
        contact,
        reference_run=table.cell(2, 3).paragraphs[0].runs[0],
    )
    _value_cell(
        table.cell(3, 1),
        data.inspection_details.strip(),
        reference_run=table.cell(3, 0).paragraphs[0].runs[0],
    )
    _value_cell(
        table.cell(4, 1),
        data.remarks.strip(),
        reference_run=table.cell(4, 0).paragraphs[0].runs[0],
    )

    if document.paragraphs[6].text or document.paragraphs[8].text:
        raise DocumentGenerationError("记录表签名值区域必须保持空白")
    result = _save_document(document)
    from .inspection import assert_document_structure

    assert_document_structure(result, expected_type="record")
    return result


def _normalize_image(item: PhotoItem) -> tuple[bytes, int, int]:
    if len(item.image) > MAX_IMAGE_BYTES:
        raise DocumentGenerationError("照片文件大小超出限制")
    try:
        with Image.open(BytesIO(item.image)) as opened:
            if opened.format not in {"JPEG", "PNG"}:
                raise DocumentGenerationError("照片文书只支持 JPEG 或 PNG 图片")
            if opened.width * opened.height > MAX_IMAGE_PIXELS:
                raise DocumentGenerationError("照片像素数量超出限制")
            opened.load()
            normalized = ImageOps.exif_transpose(opened)
            if normalized.mode not in {"RGB", "RGBA"}:
                normalized = normalized.convert("RGB")
            output = BytesIO()
            normalized.save(output, format="PNG", optimize=True)
            return output.getvalue(), normalized.width, normalized.height
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        if isinstance(exc, DocumentGenerationError):
            raise
        raise DocumentGenerationError("照片已损坏或无法解析") from exc


def _photo_size(width: int, height: int) -> tuple[int, int]:
    scale = min(PHOTO_MAX_WIDTH / width, PHOTO_MAX_HEIGHT / height)
    return int(width * scale), int(height * scale)


def _fresh_paragraph_ids(elements: list[object], block_index: int) -> None:
    para_id_name = f"{{{W14_NS}}}paraId"
    text_id_name = f"{{{W14_NS}}}textId"
    for paragraph_index, element in enumerate(elements):
        if getattr(element, "tag", None) != qn("w:p"):
            continue
        seed = f"photo-block-{block_index}-{paragraph_index}".encode()
        element.set(para_id_name, sha256(seed).hexdigest()[:8].upper())
        if element.get(text_id_name) is not None:
            element.set(text_id_name, sha256(seed + b"-text").hexdigest()[:8].upper())


def _photo_time(value: object) -> str:
    if hasattr(value, "year") and hasattr(value, "hour"):
        return format_chinese_datetime(value)  # type: ignore[arg-type]
    return str(value).strip()


def generate_photo_docx(source: DocxSource, data: PhotoDocumentData) -> bytes:
    document, _ = _template_document(source, TemplateDocumentType.PHOTO)
    body = document._element.body
    original_elements = [element for element in list(body) if element.tag != qn("w:sectPr")]
    if len(original_elements) != 9:
        raise DocumentGenerationError("照片模板单页母版结构与已确认映射不一致")
    original_text = "".join("".join(element.itertext()) for element in original_elements)
    if "[此处插入照片]" not in original_text:
        raise DocumentGenerationError("照片模板缺少图片占位锚点")
    for element in original_elements:
        body.remove(element)

    blocks: list[list[object]] = []
    for block_index in range(len(data.items)):
        elements = [deepcopy(element) for element in original_elements]
        _fresh_paragraph_ids(elements, block_index)
        for element in elements:
            body.sectPr.addprevious(element)
        blocks.append(elements)

    shared_values = (
        data.shared.shooting_location.strip(),
        data.shared.photographer.strip(),
        _photo_time(data.shared.shooting_time),
        data.shared.confirmer.strip(),
    )
    for block_index, (elements, item) in enumerate(zip(blocks, data.items, strict=True)):
        paragraphs = [Paragraph(element, document._body) for element in elements]
        if len(paragraphs) != 9:
            raise DocumentGenerationError("照片模板母版段落数量发生变化")

        image_properties = _run_properties(paragraphs[0].runs[0])
        _clear_paragraph(paragraphs[0])
        image_run = _add_run(paragraphs[0], "", properties=image_properties, underline=False)
        image_bytes, pixel_width, pixel_height = _normalize_image(item)
        image_width, image_height = _photo_size(pixel_width, pixel_height)
        image_run.add_picture(
            BytesIO(image_bytes),
            width=Emu(image_width),
            height=Emu(image_height),
        )

        field_values = (
            ("反映内容：", item.reflected_content.strip()),
            ("拍摄地点：", shared_values[0]),
            ("摄制人：", shared_values[1]),
            ("拍摄时间：", shared_values[2]),
            ("说明：", item.description.strip()),
            ("确认人：", shared_values[3]),
        )
        for paragraph, (label, value) in zip(paragraphs[2:8], field_values, strict=True):
            label_properties = _run_properties(paragraph.runs[0])
            _replace_paragraph(
                paragraph,
                (
                    (label, False, label_properties),
                    (value, True, label_properties),
                ),
            )
        if block_index < len(blocks) - 1:
            break_properties = _run_properties(paragraphs[8].runs[0])
            _clear_paragraph(paragraphs[8])
            page_break_run = _add_run(
                paragraphs[8], "", properties=break_properties, underline=False
            )
            page_break_run.add_break(WD_BREAK.PAGE)
        else:
            # The source's final paragraph contains a manual line break.  Once
            # a real image fills the reserved area that break can create an
            # otherwise blank trailing page, so retain the paragraph but clear
            # only that break.
            _clear_paragraph(paragraphs[8])

    result = _save_document(document)
    from .inspection import assert_document_structure

    assert_document_structure(
        result,
        expected_type="photo",
        expected_item_count=len(data.items),
    )
    return result


def _paragraph_styles(paragraph: Paragraph) -> list[object | None]:
    return [_run_properties(run) for run in paragraph.runs]


def _style_at(styles: list[object | None], index: int) -> object | None:
    if not styles:
        return None
    return styles[min(index, len(styles) - 1)]


def _set_update_fields(document: Document) -> None:
    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def _qa_paragraph_ids(paragraph: Paragraph, index: int) -> None:
    paragraph._p.set(
        f"{{{W14_NS}}}paraId",
        sha256(f"interview-qa-{index}".encode()).hexdigest()[:8].upper(),
    )
    if paragraph._p.get(f"{{{W14_NS}}}textId") is not None:
        paragraph._p.set(
            f"{{{W14_NS}}}textId",
            sha256(f"interview-qa-text-{index}".encode()).hexdigest()[:8].upper(),
        )


def generate_interview_docx(source: DocxSource, data: InterviewDocumentData) -> bytes:
    document, _ = _template_document(source, TemplateDocumentType.INTERVIEW)
    if len(document.paragraphs) != 12 or document.tables:
        raise DocumentGenerationError("询问笔录模板结构与已确认映射不一致")

    paragraphs = document.paragraphs
    styles = [_paragraph_styles(paragraph) for paragraph in paragraphs]

    _replace_paragraph(
        paragraphs[0],
        (
            ("第 ", False, _style_at(styles[0], 3)),
            (str(data.interview_number), True, _style_at(styles[0], 4)),
            (" 次", False, _style_at(styles[0], 5)),
        ),
    )
    _replace_paragraph(
        paragraphs[2],
        (
            ("时间：", False, _style_at(styles[2], 0)),
            (format_chinese_datetime(data.started_at), True, _style_at(styles[2], 1)),
            (" 至 ", False, _style_at(styles[2], 2)),
            (format_chinese_datetime(data.ended_at), True, _style_at(styles[2], 1)),
        ),
    )
    _replace_paragraph(
        paragraphs[3],
        (
            ("地点：", False, _style_at(styles[3], 0)),
            (data.location.strip(), True, _style_at(styles[3], 0)),
        ),
    )
    # The source labels these positions as signatures.  They remain blank for
    # handwritten signing; only the work-unit values are generated.
    _replace_paragraph(
        paragraphs[4],
        (
            ("询问人（签名）：", False, _style_at(styles[4], 0)),
            ("  工作单位：", False, _style_at(styles[4], 2)),
            (data.interviewer_unit.strip(), True, _style_at(styles[4], 1)),
        ),
    )
    _replace_paragraph(
        paragraphs[5],
        (
            ("记录人（签名）：", False, _style_at(styles[5], 0)),
            ("  工作单位：", False, _style_at(styles[5], 3)),
            (data.recorder_unit.strip(), True, _style_at(styles[5], 4)),
        ),
    )
    _replace_paragraph(
        paragraphs[6],
        (
            ("被询问人姓名：", False, _style_at(styles[6], 0)),
            (data.interviewee_name.strip(), True, _style_at(styles[6], 1)),
            ("  性别：", False, _style_at(styles[6], 3)),
            (data.gender.strip(), True, _style_at(styles[6], 4)),
            ("  年龄：", False, _style_at(styles[6], 5)),
            (data.age.strip(), True, _style_at(styles[6], 6)),
            ("  出生日期：", False, _style_at(styles[6], 7)),
            (data.birth_date.strip(), True, _style_at(styles[6], 7)),
        ),
    )
    _replace_paragraph(
        paragraphs[7],
        (
            ("证件类型：", False, _style_at(styles[7], 0)),
            (data.id_type.strip(), True, _style_at(styles[7], 1)),
            ("  证件号码：", False, _style_at(styles[7], 2)),
            (data.id_number.strip(), True, _style_at(styles[7], 1)),
        ),
    )
    _replace_paragraph(
        paragraphs[8],
        (
            ("户籍所在地：", False, _style_at(styles[8], 0)),
            (data.registered_address.strip(), True, _style_at(styles[8], 0)),
        ),
    )
    _replace_paragraph(
        paragraphs[9],
        (
            ("现住址：", False, _style_at(styles[9], 0)),
            (data.current_address.strip(), True, _style_at(styles[9], 0)),
        ),
    )
    _replace_paragraph(
        paragraphs[10],
        (
            ("联系方式：", False, _style_at(styles[10], 0)),
            (data.contact.strip(), True, _style_at(styles[10], 0)),
        ),
    )

    anchor = paragraphs[11]
    anchor_template = deepcopy(anchor._p)
    body = document._element.body
    previous_element = anchor._p
    for index, qa_item in enumerate(data.qa_items):
        if index == 0:
            paragraph = anchor
        else:
            element = deepcopy(anchor_template)
            previous_element.addnext(element)
            previous_element = element
            paragraph = Paragraph(element, document._body)
        _clear_paragraph(paragraph)
        _qa_paragraph_ids(paragraph, index)
        _add_run(
            paragraph,
            f"{qa_item.role.prefix}{qa_item.text.strip()}",
            properties=_style_at(styles[11], 0),
            underline=True,
        )
    if body.sectPr is None:
        raise DocumentGenerationError("询问笔录模板缺少页面设置")

    _set_update_fields(document)
    result = finalize_interview_page_count(_save_document(document), page_count=1)
    from .inspection import assert_document_structure

    assert_document_structure(
        result,
        expected_type="interview",
        expected_item_count=len(data.qa_items),
    )
    return result
