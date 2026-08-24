"""Coordinate-aware PDF parsing for text-based Chinese legal documents."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from uuid import UUID

from .models import KnowledgeClause

ARTICLE_START = re.compile(r"^第([零〇一二三四五六七八九十百]+)条(?:\s*)")
CHAPTER_LINE = re.compile(r"^第[零〇一二三四五六七八九十百]+章")
FOOTER_LINE = re.compile(r"^[—\-－]?\s*[０-９0-9]+\s*[—\-－]?$")
NUMERAL_DIGITS = {
    "零": 0,
    "〇": 0,
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
}


class KnowledgeParseError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class PositionedLine:
    text: str
    page_number: int
    top: float
    column: int


def chinese_numeral_to_int(value: str) -> int:
    if "百" in value:
        left, right = value.split("百", 1)
        hundreds = NUMERAL_DIGITS.get(left, 1) if left else 1
        return hundreds * 100 + (chinese_numeral_to_int(right) if right else 0)
    if "十" in value:
        left, right = value.split("十", 1)
        tens = NUMERAL_DIGITS.get(left, 1) if left else 1
        ones = NUMERAL_DIGITS.get(right, 0) if right else 0
        return tens * 10 + ones
    if len(value) > 1:
        result = 0
        for char in value:
            result = result * 10 + NUMERAL_DIGITS[char]
        return result
    return NUMERAL_DIGITS[value]


def _group_column_words(
    words: list[dict[str, object]], *, tolerance: float = 2.5
) -> list[tuple[float, str]]:
    rows: list[list[dict[str, object]]] = []
    for word in sorted(words, key=lambda item: (float(item["top"]), float(item["x0"]))):
        top = float(word["top"])
        if not rows or abs(float(rows[-1][0]["top"]) - top) > tolerance:
            rows.append([word])
        else:
            rows[-1].append(word)
    lines: list[tuple[float, str]] = []
    for row in rows:
        row.sort(key=lambda item: float(item["x0"]))
        text = "".join(str(item["text"]) for item in row).strip()
        if text:
            lines.append((float(row[0]["top"]), text))
    return lines


def extract_two_column_lines(path: Path) -> list[PositionedLine]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - deployment preflight covers it
        raise KnowledgeParseError("缺少 PDF 文本解析组件") from exc

    lines: list[PositionedLine] = []
    try:
        pdf = pdfplumber.open(path)
    except Exception as exc:
        raise KnowledgeParseError("PDF 已损坏或无法解析") from exc
    with pdf:
        if not pdf.pages:
            raise KnowledgeParseError("PDF 没有页面")
        for page_index, page in enumerate(pdf.pages):
            midpoint = float(page.width) / 2
            page_words = page.extract_words(
                x_tolerance=2,
                y_tolerance=2,
                keep_blank_chars=False,
                use_text_flow=False,
            )
            columns: tuple[list[dict[str, object]], list[dict[str, object]]] = ([], [])
            for word in page_words:
                top = float(word["top"])
                # Page furniture is not legal content.  The first page has a
                # full-width publication preamble above the two-column body.
                if top < 35 or top > float(page.height) - 24:
                    continue
                if page_index == 0 and top < 250:
                    continue
                column = 0 if float(word["x0"]) < midpoint else 1
                columns[column].append(word)
            for column_index, column_words in enumerate(columns):
                for top, text in _group_column_words(column_words):
                    normalized = unicodedata.normalize("NFKC", text).strip()
                    if not normalized or FOOTER_LINE.fullmatch(normalized):
                        continue
                    lines.append(
                        PositionedLine(
                            text=normalized,
                            page_number=page_index + 1,
                            top=top,
                            column=column_index,
                        )
                    )
    return lines


def clauses_from_lines(lines: list[PositionedLine], document_id: UUID) -> list[KnowledgeClause]:
    clauses: list[KnowledgeClause] = []
    current_identifier = ""
    current_sequence = 0
    current_parts: list[str] = []
    current_page_start = 0
    current_page_end = 0
    expected_sequence = 1

    def finish() -> None:
        nonlocal current_identifier, current_sequence, current_parts
        if not current_identifier:
            return
        content = "".join(current_parts).strip()
        if not content:
            raise KnowledgeParseError(f"{current_identifier}没有正文")
        clauses.append(
            KnowledgeClause(
                document_id=document_id,
                sequence=current_sequence,
                clause_identifier=current_identifier,
                content=content,
                page_start=current_page_start,
                page_end=current_page_end,
            )
        )
        current_identifier = ""
        current_parts = []

    for line in lines:
        match = ARTICLE_START.match(line.text)
        if match:
            sequence = chinese_numeral_to_int(match.group(1))
            # A citation can begin a wrapped line.  Only the next expected
            # article starts a new record, which prevents false splitting.
            if sequence == expected_sequence:
                finish()
                current_sequence = sequence
                current_identifier = match.group(0).strip()
                current_parts = [line.text[match.end() :]]
                current_page_start = line.page_number
                current_page_end = line.page_number
                expected_sequence += 1
                continue
        if current_identifier:
            if CHAPTER_LINE.match(line.text):
                continue
            current_parts.append(line.text)
            current_page_end = line.page_number
    finish()
    return clauses


def parse_text_pdf(path: Path, document_id: UUID) -> list[KnowledgeClause]:
    if path.suffix.lower() != ".pdf":
        raise KnowledgeParseError("仅支持文本型 PDF")
    clauses = clauses_from_lines(extract_two_column_lines(path), document_id)
    if not clauses:
        raise KnowledgeParseError("未识别到条款，扫描件需要人工或 OCR 处理")
    return clauses


def parse_text_docx(path: Path, document_id: UUID) -> list[KnowledgeClause]:
    if path.suffix.lower() != ".docx":
        raise KnowledgeParseError("仅支持 DOCX")
    try:
        from docx import Document

        document = Document(path)
    except Exception as exc:
        raise KnowledgeParseError("DOCX 已损坏或无法解析") from exc
    texts: list[str] = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text.strip() for cell in row.cells)
    lines = [
        PositionedLine(text=unicodedata.normalize("NFKC", text), page_number=1, top=index, column=0)
        for index, text in enumerate(texts)
        if text
    ]
    clauses = clauses_from_lines(lines, document_id)
    if not clauses:
        raise KnowledgeParseError("未识别到条款，请人工补充分段或映射")
    return clauses


def parse_fire_law_pdf(path: Path, document_id: UUID) -> list[KnowledgeClause]:
    clauses = parse_text_pdf(path, document_id)
    identifiers = [clause.sequence for clause in clauses]
    if identifiers != list(range(1, 75)):
        missing = sorted(set(range(1, 75)) - set(identifiers))
        raise KnowledgeParseError(f"消防法条款顺序校验失败，缺少条款: {missing}")
    return clauses
