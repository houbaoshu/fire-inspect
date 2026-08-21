from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT_DIR = Path("templates/sample")
BODY_FONT = "Arial Unicode MS"
TITLE_FONT = "Arial Unicode MS"
RED = "9B1C1C"
MUTED = "6B625D"
LIGHT_FILL = "F2F4F7"
WIDTH_DXA = 9360


def set_run(run, *, size=11, bold=False, color="000000", font=BODY_FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def write_cell(cell, text, *, label=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.1
    set_run(paragraph.add_run(text), bold=label, color=RED if label else "000000")
    if label:
        shade(cell, LIGHT_FILL)


def setup_document(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run("消防监督检查文书技术样例"), size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("系统生成初稿 · 请由执法人员核对确认"), size=8.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(title), size=20, bold=True, color=RED, font=TITLE_FONT)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    set_run(subtitle.add_run("MVP V1 可填充技术样例"), size=9.5, color=MUTED)
    return doc


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size=13, bold=True, color=RED)


def add_metadata(doc, rows):
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    set_table_geometry(table, [1450, 3230, 1450, 3230])
    for row, values in zip(table.rows, rows, strict=True):
        for index, value in enumerate(values):
            write_cell(row.cells[index], value, label=index % 2 == 0)


def add_content_box(doc, placeholder, min_lines=3):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [WIDTH_DXA])
    cell = table.cell(0, 0)
    write_cell(cell, placeholder)
    for _ in range(min_lines - 1):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(8)


def inspection_record():
    doc = setup_document("消防监督检查记录")
    add_metadata(
        doc,
        [
            ("检查名称", "{{ inspection_title }}", "检查时间", "{{ inspection_date }}"),
            ("被检查单位", "{{ organization }}", "检查人员", "{{ inspectors }}"),
            ("单位地址", "{{ address }}", "记录状态", "人工确认后生效"),
        ],
    )
    add_heading(doc, "一、检查事实")
    add_content_box(doc, "{{ facts }}", 5)
    add_heading(doc, "二、相关法规依据（供执法参考）")
    add_content_box(doc, "{{ legal_basis }}", 4)
    return doc


def photo_record():
    doc = setup_document("消防监督检查照片记录")
    add_metadata(
        doc,
        [
            ("检查名称", "{{ inspection_title }}", "检查时间", "{{ inspection_date }}"),
            ("被检查单位", "{{ organization }}", "单位地址", "{{ address }}"),
        ],
    )
    add_heading(doc, "证据照片")
    photo = doc.add_table(rows=1, cols=1)
    photo.style = "Table Grid"
    set_table_geometry(photo, [WIDTH_DXA])
    write_cell(photo.cell(0, 0), "{{ photo }}", align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(7):
        photo.cell(0, 0).add_paragraph()
    add_heading(doc, "照片说明")
    add_content_box(doc, "{{ photo_caption }}", 4)
    return doc


def inquiry_record():
    doc = setup_document("询问笔录")
    add_metadata(
        doc,
        [
            ("检查名称", "{{ inspection_title }}", "询问时间", "{{ inspection_date }}"),
            ("被检查单位", "{{ organization }}", "询问地点", "{{ address }}"),
            ("询问人员", "{{ inspectors }}", "笔录状态", "人工确认后生效"),
        ],
    )
    add_heading(doc, "询问与回答")
    add_content_box(doc, "{{ questions_answers }}", 12)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    set_run(
        note.add_run("说明：本笔录由系统依据已确认对话事实生成，须经人工核对、修改和确认。"),
        size=9.5,
        color=MUTED,
    )
    return doc


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, document in (
        ("inspection-record.docx", inspection_record()),
        ("photo-record.docx", photo_record()),
        ("inquiry-record.docx", inquiry_record()),
    ):
        document.core_properties.title = name.removesuffix(".docx")
        document.core_properties.author = "消防监督检查 AI 辅助系统"
        document.save(OUTPUT_DIR / name)


if __name__ == "__main__":
    main()
