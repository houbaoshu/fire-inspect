from __future__ import annotations

from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from fire_inspect.documents import (
    InspectionType,
    InterviewDocumentData,
    PhotoDocumentData,
    PhotoItem,
    PhotoSharedFields,
    QaItem,
    QaRole,
    RecordDocumentData,
    assert_document_structure,
    finalize_interview_page_count,
    generate_interview_docx,
    generate_photo_docx,
    generate_record_docx,
)

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def _all_text(document: Document) -> str:
    text = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(text)


def _record_data(inspection_type: InspectionType) -> RecordDocumentData:
    return RecordDocumentData(
        document_number="桐消监记〔2026〕18号",
        inspection_date=date(2026, 8, 24),
        inspection_type=inspection_type,
        unit_name="某某商业管理有限公司",
        responsible_person="张某",
        address="浙江省某市某区某路 88 号",
        contact_name="李某",
        contact_phone="13800000000",
        inspection_details="大厅存在堆放物品，相关区域用途需要进一步核实。",
        remarks="仅记录现场情况，不作自动执法结论。",
    )


@pytest.mark.parametrize("inspection_type", list(InspectionType))
def test_record_document_fills_correct_cells_and_exactly_one_checkbox(
    template_dir: Path, inspection_type: InspectionType
) -> None:
    template = template_dir / "监督检查记录表.docx"
    before = template.read_bytes()
    result = generate_record_docx(template, _record_data(inspection_type))
    document = Document(BytesIO(result))
    table = document.tables[0]

    assert table.cell(1, 2).text == "某某商业管理有限公司"
    assert table.cell(1, 4).text == "张某"
    assert table.cell(2, 2).text == "浙江省某市某区某路 88 号"
    assert table.cell(2, 4).text == "李某 13800000000"
    assert "大厅存在堆放物品" in table.cell(3, 1).text
    assert "不作自动执法结论" in table.cell(4, 1).text
    assert sum(paragraph.text.count("☑") for paragraph in table.cell(0, 0).paragraphs) == 1
    assert sum(paragraph.text.count("□") for paragraph in table.cell(0, 0).paragraphs) == 6
    assert f"☑ {inspection_type.value}" in table.cell(0, 0).text
    assert document.paragraphs[3].text == "编号: 桐消监记〔2026〕18号"
    assert document.paragraphs[4].text == "检查时间: 2026年8月24日"
    assert all(run.underline is not True for run in document.paragraphs[3].runs)
    assert all(run.underline is not True for run in table.cell(1, 2).paragraphs[0].runs)
    assert document.paragraphs[6].text == ""
    assert document.paragraphs[8].text == ""
    assert len(table.rows) == 5 and len(table.columns) == 5
    assert sha256(template.read_bytes()).digest() == sha256(before).digest()
    assert_document_structure(result, expected_type="record")


def test_generated_documents_are_byte_deterministic(template_dir: Path) -> None:
    """Visual acceptance regenerates fixtures at review time, so bytes must not vary."""

    from fire_inspect.documents.package import FIXED_ENTRY_TIMESTAMP

    template = template_dir / "监督检查记录表.docx"
    first = generate_record_docx(template, _record_data(InspectionType.CONSTRUCTION_SITE))
    second = generate_record_docx(template, _record_data(InspectionType.CONSTRUCTION_SITE))
    assert first == second
    with ZipFile(BytesIO(first)) as archive:
        assert all(
            info.date_time == FIXED_ENTRY_TIMESTAMP for info in archive.infolist()
        )


def test_record_long_content_can_grow_without_exact_height_clipping(template_dir: Path) -> None:
    data = _record_data(InspectionType.OTHER)
    data = RecordDocumentData(
        **{
            **data.as_dict(),
            "inspection_details": "现场检查情况。" * 400,
            "remarks": "备注。" * 120,
        }
    )
    result = generate_record_docx(template_dir / "监督检查记录表.docx", data)
    document = Document(BytesIO(result))
    table = document.tables[0]

    assert table.cell(3, 1).text == "现场检查情况。" * 400
    assert table.rows[3].height_rule is not None or table.rows[3].height is not None
    row_xml = table.rows[3]._tr.xml
    assert 'w:hRule="exact"' not in row_xml
    assert_document_structure(result, expected_type="record")


def _photo_data(images: tuple[bytes, bytes, bytes]) -> PhotoDocumentData:
    return PhotoDocumentData(
        items=tuple(
            PhotoItem(
                image=image,
                reflected_content=f"第 {index} 张照片反映内容",
                description=f"第 {index} 张照片说明",
                source_name=f"photo-{index}.png",
            )
            for index, image in enumerate(images, start=1)
        ),
        shared=PhotoSharedFields(
            shooting_location="某单位一层大厅",
            photographer="检查员甲",
            shooting_time=datetime(2026, 8, 24, 10, 30),
            confirmer="",
        ),
    )


def test_three_photo_document_has_three_page_blocks_and_three_real_images(
    template_dir: Path, three_test_images: tuple[bytes, bytes, bytes]
) -> None:
    template = template_dir / "照片.docx"
    before = template.read_bytes()
    result = generate_photo_docx(template, _photo_data(three_test_images))
    document = Document(BytesIO(result))
    text = _all_text(document)

    assert len(document.inline_shapes) == 3
    assert "[此处插入照片]" not in text
    assert text.count("拍摄地点：某单位一层大厅") == 3
    assert text.count("摄制人：检查员甲") == 3
    assert text.count("拍摄时间：2026年8月24日 10:30") == 3
    for index in range(1, 4):
        assert text.count(f"第 {index} 张照片反映内容") == 1
        assert text.count(f"第 {index} 张照片说明") == 1

    page_breaks = document.element.xpath('.//w:br[@w:type="page"]')
    assert len(page_breaks) == 2

    with ZipFile(BytesIO(result)) as archive:
        image_parts = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert len(image_parts) == 3
    assert sha256(template.read_bytes()).digest() == sha256(before).digest()
    assert_document_structure(result, expected_type="photo", expected_item_count=3)


def test_photo_labels_are_not_underlined_but_values_are(
    template_dir: Path, three_test_images: tuple[bytes, bytes, bytes]
) -> None:
    result = generate_photo_docx(template_dir / "照片.docx", _photo_data(three_test_images))
    document = Document(BytesIO(result))
    field_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(
            ("反映内容：", "拍摄地点：", "摄制人：", "拍摄时间：", "说明：", "确认人：")
        )
    ]

    assert len(field_paragraphs) == 18
    for paragraph in field_paragraphs:
        assert len(paragraph.runs) >= 2
        assert paragraph.runs[0].underline is not True
        if paragraph.runs[1].text:
            assert paragraph.runs[1].underline is True


def test_photo_images_keep_pixel_aspect_ratio_and_stay_inside_region(
    template_dir: Path, three_test_images: tuple[bytes, bytes, bytes]
) -> None:
    result = generate_photo_docx(template_dir / "照片.docx", _photo_data(three_test_images))
    document = Document(BytesIO(result))
    expected_ratios = (800 / 450, 450 / 800, 1.0)

    for shape, expected_ratio in zip(document.inline_shapes, expected_ratios, strict=True):
        assert shape.width.inches <= 6.0
        assert shape.height.inches <= 5.0
        assert shape.width / shape.height == pytest.approx(expected_ratio, rel=0.002)


def _interview_data(*, long: bool = False) -> InterviewDocumentData:
    repeat = 180 if long else 1
    qa = tuple(
        item
        for index in range(repeat)
        for item in (
            QaItem(QaRole.QUESTION, f"第 {index + 1} 个问题：请说明现场情况。"),
            QaItem(QaRole.ANSWER, "现场有物品堆放，区域用途情况需要进一步核实。"),
        )
    )
    return InterviewDocumentData(
        interview_number=2,
        started_at=datetime(2026, 8, 24, 23, 55),
        ended_at=datetime(2026, 8, 25, 0, 20),
        location="某市消防救援大队询问室",
        interviewer_name="检查员甲",
        interviewer_unit="某市消防救援大队",
        recorder_name="记录员乙",
        recorder_unit="某市消防救援大队",
        interviewee_name="王某",
        gender="男",
        age="42",
        birth_date="1984年3月",
        id_type="居民身份证",
        id_number="330000********0018",
        registered_address="浙江省某市某区某街道某社区",
        current_address="浙江省某市某区某路某号某幢某室" + ("很长地址" * 30 if long else ""),
        contact="138****0000",
        qa_items=qa,
    )


def test_interview_header_and_confirmed_qa_are_underlined_and_independent(
    template_dir: Path,
) -> None:
    result = generate_interview_docx(template_dir / "询问笔录.docx", _interview_data())
    document = Document(BytesIO(result))
    text = [paragraph.text for paragraph in document.paragraphs]

    assert "时间：2026年8月24日 23:55 至 2026年8月25日 00:20" in text
    assert "地点：某市消防救援大队询问室" in text
    assert "证件类型：居民身份证  证件号码：330000********0018" in text
    qa_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(("问：", "答："))
    ]
    assert [paragraph.text[:2] for paragraph in qa_paragraphs] == ["问：", "答："]
    assert all(
        paragraph.runs and all(run.underline is True for run in paragraph.runs)
        for paragraph in qa_paragraphs
    )

    for paragraph in document.paragraphs:
        if "：" in paragraph.text and paragraph.text not in {"询 问 笔 录"}:
            value_runs = [run for run in paragraph.runs if run.underline is True]
            if paragraph.text.startswith("第") or paragraph.text.startswith(("问：", "答：")):
                assert value_runs
            else:
                assert value_runs, paragraph.text
    assert_document_structure(result, expected_type="interview", expected_item_count=2)


def test_interview_rejects_unconfirmed_or_empty_qa() -> None:
    with pytest.raises(ValueError, match="人工确认"):
        QaItem(QaRole.UNKNOWN, "尚未确认角色")
    with pytest.raises(ValueError, match="正文"):
        QaItem(QaRole.QUESTION, "  ")


def test_interview_preserves_all_footer_fields_and_requests_field_refresh(
    template_dir: Path,
) -> None:
    generated = generate_interview_docx(template_dir / "询问笔录.docx", _interview_data())
    finalized = finalize_interview_page_count(generated, page_count=3)
    with ZipFile(BytesIO(finalized)) as archive:
        footer_names = sorted(name for name in archive.namelist() if name.startswith("word/footer"))
        assert len(footer_names) == 3
        for name in footer_names:
            root = etree.fromstring(archive.read(name))
            instructions = " ".join(root.xpath(".//w:instrText/text()", namespaces=NS))
            simple = " ".join(root.xpath(".//w:fldSimple/@w:instr", namespaces=NS))
            assert "PAGE" in instructions
            assert "NUMPAGES" in simple
            cached_total = root.xpath(
                './/w:fldSimple[contains(@w:instr, "NUMPAGES")]//w:t/text()', namespaces=NS
            )
            assert cached_total == ["3"]

        settings = etree.fromstring(archive.read("word/settings.xml"))
        assert settings.xpath('boolean(.//w:updateFields[@w:val="true"])', namespaces=NS)


def test_long_interview_keeps_natural_paragraphs_without_fixed_height_structures(
    template_dir: Path,
) -> None:
    result = generate_interview_docx(
        template_dir / "询问笔录.docx",
        _interview_data(long=True),
    )
    document = Document(BytesIO(result))
    qa_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith(("问：", "答："))
    ]

    assert len(qa_paragraphs) == 360
    assert not document.tables
    assert all("w:framePr" not in paragraph._p.xml for paragraph in qa_paragraphs)
    assert_document_structure(result, expected_type="interview", expected_item_count=360)
