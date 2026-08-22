import io

from docx import Document
from PIL import Image

from backend.app.document_service import concise_photo_description, format_interview_qa


def template_bytes() -> bytes:
    document = Document()
    document.add_heading("消防监督检查记录", level=1)
    document.add_paragraph("固定文字：本页版式不得重建")
    document.add_paragraph("被检查单位：__________")
    document.add_paragraph("检查情况：__________")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def photo_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("[此处插入照片]")
    document.add_paragraph("")
    document.add_paragraph("反映内容： ")
    document.add_paragraph("拍摄地点： ")
    document.add_paragraph("摄制人： ")
    document.add_paragraph("拍摄时间： ")
    document.add_paragraph("说明： ")
    document.add_paragraph("确认人： ")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def interview_template_bytes() -> bytes:
    document = Document()
    document.add_heading("询 问 笔 录", level=1)
    document.add_paragraph("时间：                           至")
    document.add_paragraph("地点：")
    document.add_paragraph("询问人（签名）：              工作单位：")
    document.add_paragraph("记录人（签名）：     工作单位：               ")
    document.add_paragraph("被询问人姓名：     性别：  年龄：     出生日期：")
    document.add_paragraph("证件类型：                         证件号码：")
    document.add_paragraph("户籍所在地：")
    document.add_paragraph("现住址：")
    document.add_paragraph("联系方式：")
    document.add_paragraph("")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def inspection_table_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("检查时间:    年   月   日")
    table = document.add_table(rows=5, cols=2)
    for row, label in zip(
        table.rows,
        [
            "单位（场所）\n名称",
            "法定代表/\n主要负责人",
            "联系人及\n联系电话",
            "检查\n内容\n和\n情况",
            "备注",
        ],
        strict=True,
    ):
        row.cells[0].text = label
        row.cells[1].text = ""
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def image_bytes(color: str) -> bytes:
    image = Image.new("RGB", (80, 60), color)
    output = io.BytesIO()
    image.save(output, format="JPEG")
    return output.getvalue()


def upload_image(client, headers, inspection_id: str, filename: str, data: bytes) -> dict:
    upload = client.post(
        f"/api/v1/inspections/{inspection_id}/upload-sessions",
        json={"filename": filename, "content_type": "image/jpeg", "size_bytes": len(data)},
        headers=headers,
    ).json()
    part = client.put(
        f"/api/v1/upload-sessions/{upload['id']}/parts/1",
        content=data,
        headers=headers,
    ).json()
    response = client.post(
        f"/api/v1/upload-sessions/{upload['id']}/complete",
        json={"parts": [part]},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return response.json()


def upload_and_confirm_template(
    client, headers, name: str, document_type: str, filename: str, data: bytes
) -> dict:
    uploaded = client.post(
        "/api/v1/templates",
        data={"name": name, "document_type": document_type},
        files={
            "file": (
                filename,
                data,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=headers,
    ).json()
    for region in uploaded["regions"]:
        confirmed = client.patch(
            f"/api/v1/templates/regions/{region['id']}",
            json={"confirmed": True, "row_version": region["row_version"]},
            headers=headers,
        )
        assert confirmed.status_code == 200, confirmed.text
    response = client.post(
        f"/api/v1/templates/{uploaded['template_id']}/versions/{uploaded['id']}/confirm",
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return uploaded


def test_photo_and_interview_templates_recognize_fixed_label_fields(client, admin_headers):
    photo = client.post(
        "/api/v1/templates",
        data={"name": "照片模板", "document_type": "photo_document"},
        files={
            "file": (
                "photo.docx",
                photo_template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=admin_headers,
    )
    assert photo.status_code == 201, photo.text
    photo_regions = photo.json()["regions"]
    assert {region["label"] for region in photo_regions} == {
        "照片",
        "反映内容",
        "拍摄地点",
        "摄制人",
        "拍摄时间",
        "说明",
        "确认人",
    }
    photo_sources = {region["label"]: region["source_path"] for region in photo_regions}
    assert photo_sources["照片"] == "document.photo_entries"
    assert photo_sources["反映内容"] == "facts.confirmed_text"
    assert next(region for region in photo_regions if region["label"] == "照片")["repeatable"]
    assert all(region["repeatable"] for region in photo_regions)

    interview = client.post(
        "/api/v1/templates",
        data={"name": "询问笔录模板", "document_type": "interview_transcript"},
        files={
            "file": (
                "interview.docx",
                interview_template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=admin_headers,
    )
    assert interview.status_code == 201, interview.text
    interview_regions = interview.json()["regions"]
    labels = [region["label"] for region in interview_regions]
    assert "被询问人姓名" in labels
    assert "证件号码" in labels
    assert "询问问答正文" in labels
    assert sum(label == "工作单位" for label in labels) == 2
    assert (
        next(region for region in interview_regions if region["label"] == "询问问答正文")[
            "source_path"
        ]
        == "transcript.confirmed_text"
    )


def test_inspection_table_fields_have_stable_semantic_sources(client, admin_headers):
    uploaded = client.post(
        "/api/v1/templates",
        data={"name": "监督检查记录表", "document_type": "inspection_record"},
        files={
            "file": (
                "inspection.docx",
                inspection_table_template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=admin_headers,
    )
    assert uploaded.status_code == 201, uploaded.text
    regions = uploaded.json()["regions"]
    sources = {region["label"].replace("\n", ""): region["source_path"] for region in regions}
    assert sources["法定代表/主要负责人"] == "custom.legal_representative"
    assert sources["联系人及联系电话"] == "custom.contact_and_phone"
    assert sources["检查内容和情况"] == "facts.confirmed_text"
    assert sources["备注"] == "custom.remarks"
    assert all(
        "custom." not in source or len(source.rsplit(".", 1)[-1]) != 12
        for source in sources.values()
    )


def test_photo_document_repeats_complete_block_for_each_selected_image(client, admin_headers):
    uploaded = client.post(
        "/api/v1/templates",
        data={"name": "照片模板", "document_type": "photo_document"},
        files={
            "file": (
                "photo.docx",
                photo_template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=admin_headers,
    ).json()
    for region in uploaded["regions"]:
        response = client.patch(
            f"/api/v1/templates/regions/{region['id']}",
            json={"confirmed": True, "row_version": region["row_version"]},
            headers=admin_headers,
        )
        assert response.status_code == 200, response.text
    assert (
        client.post(
            f"/api/v1/templates/{uploaded['template_id']}/versions/{uploaded['id']}/confirm",
            headers=admin_headers,
        ).status_code
        == 200
    )
    inspection = client.post(
        "/api/v1/inspections",
        json={
            "title": "照片文书测试",
            "address": "一层疏散通道",
            "dynamic_attributes": {
                "photographer": "张三",
                "photo_confirmer": "李四",
            },
        },
        headers=admin_headers,
    ).json()
    upload_image(client, admin_headers, inspection["id"], "第一张.jpg", image_bytes("red"))
    upload_image(client, admin_headers, inspection["id"], "第二张.jpg", image_bytes("blue"))
    evidence = client.get(
        f"/api/v1/inspections/{inspection['id']}/evidence", headers=admin_headers
    ).json()
    selected = [item["id"] for item in evidence if item["type"] == "image"]
    draft = client.post(
        f"/api/v1/inspections/{inspection['id']}/documents",
        json={
            "template_version_id": uploaded["id"],
            "title": "照片文书初稿",
            "selected_evidence_ids": selected,
        },
        headers=admin_headers,
    )
    assert draft.status_code == 200, draft.text
    description = next(field for field in draft.json()["fields"] if field["label"] == "说明")
    assert description["field_type"] == "multiline"
    assert len(description["value"]) == 2
    per_photo_fields = {"照片", "反映内容", "说明"}
    common_fields = {"拍摄地点", "摄制人", "拍摄时间", "确认人"}
    for field in draft.json()["fields"]:
        if field["label"] in per_photo_fields:
            assert isinstance(field["value"], list)
            assert len(field["value"]) == 2
        if field["label"] in common_fields:
            assert not isinstance(field["value"], list)
    edited = client.patch(
        f"/api/v1/document-fields/{description['id']}",
        json={
            "value": ["第一张独立说明", "第二张独立说明"],
            "row_version": description["row_version"],
        },
        headers=admin_headers,
    )
    assert edited.status_code == 200, edited.text
    exported = client.post(f"/api/v1/documents/{draft.json()['id']}/export", headers=admin_headers)
    assert exported.status_code == 200, exported.text
    document = Document(io.BytesIO(exported.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert len(document.inline_shapes) == 2
    assert text.count("反映内容：") == 2
    assert text.count("拍摄地点：一层疏散通道") == 2
    assert text.count("摄制人：张三") == 2
    assert text.count("确认人：李四") == 2
    assert "说明：第一张独立说明" in text
    assert "说明：第二张独立说明" in text


def test_photo_description_keeps_only_one_short_objective_sentence():
    description = (
        "画面中可见疏散通道内堆放纸箱，通道宽度受到影响。"
        "第二句包含不应进入照片说明的冗长分析和推断。"
    )
    assert concise_photo_description(description) == (
        "画面中可见疏散通道内堆放纸箱，通道宽度受到影响。"
    )
    assert concise_photo_description("甲" * 80) == f"{'甲' * 59}…"
    assert concise_photo_description("", "关键帧 1") == "关键帧 1"
    assert (
        concise_photo_description(
            "画面中可见以下明确的消防检查相关内容：\n"
            "1. 墙面安装有三类消防报警/控制设备：\n"
            "- 设备细节很长。"
        )
        == "墙面安装有三类消防报警/控制设备。"
    )


def test_three_document_types_share_inspection_fields_and_evidence_sources(client, admin_headers):
    photo_template = upload_and_confirm_template(
        client,
        admin_headers,
        "共享照片模板",
        "photo_document",
        "photo.docx",
        photo_template_bytes(),
    )
    interview_template = upload_and_confirm_template(
        client,
        admin_headers,
        "共享询问模板",
        "interview_transcript",
        "interview.docx",
        interview_template_bytes(),
    )
    inspection = client.post(
        "/api/v1/inspections",
        json={"title": "共享证据链测试", "address": "原检查地点"},
        headers=admin_headers,
    ).json()
    upload_image(
        client,
        admin_headers,
        inspection["id"],
        "共享关键帧.jpg",
        image_bytes("green"),
    )
    evidence = client.get(
        f"/api/v1/inspections/{inspection['id']}/evidence", headers=admin_headers
    ).json()
    photo = client.post(
        f"/api/v1/inspections/{inspection['id']}/documents",
        json={
            "template_version_id": photo_template["id"],
            "title": "照片文书",
            "selected_evidence_ids": [evidence[0]["id"]],
        },
        headers=admin_headers,
    ).json()
    interview = client.post(
        f"/api/v1/inspections/{inspection['id']}/documents",
        json={
            "template_version_id": interview_template["id"],
            "title": "询问笔录",
        },
        headers=admin_headers,
    ).json()
    confirmed = client.post(f"/api/v1/documents/{interview['id']}/confirm", headers=admin_headers)
    assert confirmed.status_code == 200, confirmed.text

    photo_location = next(field for field in photo["fields"] if field["label"] == "拍摄地点")
    assert photo_location["is_shared"]
    assert photo_location["source_path"] == "inspection.address"
    assert next(field for field in photo["fields"] if field["label"] == "反映内容")[
        "is_evidence_linked"
    ]
    changed = client.patch(
        f"/api/v1/document-fields/{photo_location['id']}",
        json={"value": "同步后的检查地点", "row_version": photo_location["row_version"]},
        headers=admin_headers,
    )
    assert changed.status_code == 200, changed.text

    refreshed_inspection = client.get(
        f"/api/v1/inspections/{inspection['id']}", headers=admin_headers
    ).json()
    assert refreshed_inspection["address"] == "同步后的检查地点"
    refreshed_interview = client.get(
        f"/api/v1/documents/{interview['id']}", headers=admin_headers
    ).json()
    interview_location = next(
        field for field in refreshed_interview["fields"] if field["label"] == "地点"
    )
    assert refreshed_interview["status"] == "draft"
    assert interview_location["value"] == "同步后的检查地点"
    assert interview_location["source"] == "human"
    assert interview_location["is_locked"]
    assert interview_location["is_shared"]
    assert next(
        field for field in refreshed_interview["fields"] if field["label"] == "询问问答正文"
    )["is_evidence_linked"]

    inspection_update = client.patch(
        f"/api/v1/inspections/{inspection['id']}",
        json={
            "address": "检查主数据更新后的地点",
            "row_version": refreshed_inspection["row_version"],
        },
        headers=admin_headers,
    )
    assert inspection_update.status_code == 200, inspection_update.text
    for document_id, label in (
        (photo["id"], "拍摄地点"),
        (interview["id"], "地点"),
    ):
        refreshed = client.get(f"/api/v1/documents/{document_id}", headers=admin_headers).json()
        assert (
            next(field for field in refreshed["fields"] if field["label"] == label)["value"]
            == "检查主数据更新后的地点"
        )


def test_interview_transcript_is_formatted_as_questions_and_answers():
    assert format_interview_qa(["你叫什么名字？张三。现场是否有人值班？没有。"]) == (
        "问：你叫什么名字？\n答：张三。\n问：现场是否有人值班？\n答：没有。"
    )
    assert format_interview_qa(["仅有一段尚未区分说话人的转写。"]) == (
        "问：\n答：仅有一段尚未区分说话人的转写。"
    )


def test_interview_multi_field_paragraphs_fill_without_removing_labels(client, admin_headers):
    uploaded = client.post(
        "/api/v1/templates",
        data={"name": "询问笔录模板", "document_type": "interview_transcript"},
        files={
            "file": (
                "interview.docx",
                interview_template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=admin_headers,
    ).json()
    for region in uploaded["regions"]:
        confirmed = client.patch(
            f"/api/v1/templates/regions/{region['id']}",
            json={"confirmed": True, "row_version": region["row_version"]},
            headers=admin_headers,
        )
        assert confirmed.status_code == 200, confirmed.text
    confirmed_mapping = client.post(
        f"/api/v1/templates/{uploaded['template_id']}/versions/{uploaded['id']}/confirm",
        headers=admin_headers,
    )
    assert confirmed_mapping.status_code == 200, confirmed_mapping.text

    inspection = client.post(
        "/api/v1/inspections",
        json={
            "title": "询问测试",
            "address": "一层消防控制室",
            "dynamic_attributes": {
                "interview_time_range": "2026年8月21日9时至10时",
                "questioner_name": "李检查员",
                "questioner_work_unit": "消防救援大队",
                "recorder_name": "王检查员",
                "recorder_work_unit": "消防救援大队",
                "interviewee_name": "张三",
                "interviewee_gender": "男",
                "interviewee_age": "35",
                "interviewee_birth_date": "1991年1月1日",
            },
        },
        headers=admin_headers,
    ).json()
    draft = client.post(
        f"/api/v1/inspections/{inspection['id']}/documents",
        json={"template_version_id": uploaded["id"], "title": "询问笔录初稿"},
        headers=admin_headers,
    )
    assert draft.status_code == 200, draft.text
    exported = client.post(f"/api/v1/documents/{draft.json()['id']}/export", headers=admin_headers)
    assert exported.status_code == 200, exported.text
    output = Document(io.BytesIO(exported.content))
    text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert "地点：一层消防控制室" in text
    assert "询问人（签名）：李检查员工作单位：消防救援大队" in text
    assert "被询问人姓名：张三性别：男年龄：35出生日期：1991年1月1日" in text


def test_template_mapping_document_edit_and_export_preserve_fixed_content(client, admin_headers):
    uploaded = client.post(
        "/api/v1/templates",
        data={"name": "检查记录模板", "document_type": "inspection_record"},
        files={
            "file": (
                "record.docx",
                template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=admin_headers,
    )
    assert uploaded.status_code == 201, uploaded.text
    version = uploaded.json()
    assert version["status"] == "succeeded"
    assert len(version["regions"]) == 2
    for region in version["regions"]:
        response = client.patch(
            f"/api/v1/templates/regions/{region['id']}",
            json={"confirmed": True, "row_version": region["row_version"]},
            headers=admin_headers,
        )
        assert response.status_code == 200
    confirmed = client.post(
        f"/api/v1/templates/{version['template_id']}/versions/{version['id']}/confirm",
        headers=admin_headers,
    )
    assert confirmed.status_code == 200

    inspection = client.post(
        "/api/v1/inspections",
        json={"title": "文书测试", "unit_name": "测试单位"},
        headers=admin_headers,
    ).json()
    fact = client.post(
        f"/api/v1/inspections/{inspection['id']}/facts",
        json={"title": "消防控制室无人值班", "evidence_ids": []},
        headers=admin_headers,
    ).json()
    confirmed_fact = client.post(
        f"/api/v1/facts/{fact['id']}/confirm", headers=admin_headers
    ).json()
    draft = client.post(
        f"/api/v1/inspections/{inspection['id']}/documents",
        json={"template_version_id": version["id"], "title": "检查记录初稿"},
        headers=admin_headers,
    )
    assert draft.status_code == 200, draft.text
    draft_data = draft.json()
    fact_field = next(
        field for field in draft_data["fields"] if field["field_key"] == "facts.confirmed_text"
    )
    edited = client.patch(
        f"/api/v1/document-fields/{fact_field['id']}",
        json={"value": "人工核对后的检查情况", "row_version": fact_field["row_version"]},
        headers=admin_headers,
    )
    assert edited.status_code == 200
    edited_field = next(
        field for field in edited.json()["fields"] if field["id"] == fact_field["id"]
    )
    assert edited_field["source"] == "human"
    assert edited_field["is_locked"] is True

    changed_fact = client.patch(
        f"/api/v1/facts/{fact['id']}",
        json={
            "title": "人工修正后的消防控制室事实",
            "row_version": confirmed_fact["row_version"],
        },
        headers=admin_headers,
    )
    assert changed_fact.status_code == 200
    refreshed = client.get(f"/api/v1/documents/{draft_data['id']}").json()
    stale_field = next(field for field in refreshed["fields"] if field["id"] == fact_field["id"])
    assert stale_field["is_stale"] is True
    assert stale_field["value"] == "人工核对后的检查情况"

    proposal = client.post(
        f"/api/v1/documents/{draft_data['id']}/regenerate",
        json={"field_ids": [fact_field["id"]]},
        headers=admin_headers,
    ).json()[0]
    assert proposal["locked"] is True
    assert "人工修正后的消防控制室事实" in proposal["proposed_value"]

    exported = client.post(f"/api/v1/documents/{draft_data['id']}/export", headers=admin_headers)
    assert exported.status_code == 200, exported.text
    output = Document(io.BytesIO(exported.content))
    text = "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert "固定文字：本页版式不得重建" in text
    assert "测试单位" in text
    assert "人工核对后的检查情况" in text
