from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from fire_inspect.jobs.management.commands.runworker import Command as WorkerCommand
from fire_inspect.jobs.models import ProcessingJob
from fire_inspect.jobs.services import acquire_job_by_id
from fire_inspect.knowledge.django_models import (
    KnowledgeClauseRecord,
    KnowledgeDocumentRecord,
)

ROOT = Path(__file__).resolve().parents[2]
PDF = ROOT / "testdata/knowledge/reg_ec997c582c9c29a6.pdf"
PDF_SHA256 = "92580dffac57f26f18955271332f819841ba4ad968964626d93146f3094b2576"


def uploaded_fire_law() -> SimpleUploadedFile:
    return SimpleUploadedFile(PDF.name, PDF.read_bytes(), content_type="application/pdf")


def upload_as_admin(
    authenticated_client,
    admin_user,
    *,
    idempotency_key: str = "upload-fire-law",
):
    client = authenticated_client(admin_user)
    response = client.post(
        "/api/v1/knowledge/documents/",
        {
            "file": uploaded_fire_law(),
            "title": "中华人民共和国消防法（测试样本）",
            "category": "law",
        },
        format="multipart",
        HTTP_IDEMPOTENCY_KEY=idempotency_key,
    )
    assert response.status_code == 202, response.data
    return response


def run_parse_job(*, document_id) -> ProcessingJob:
    job = ProcessingJob.objects.get(
        task_type=ProcessingJob.TaskType.KNOWLEDGE_PARSE,
        scope_type="knowledge_document",
        scope_id=document_id,
    )
    claimed = acquire_job_by_id(job_id=job.id, worker_id="knowledge-test")
    assert claimed is not None
    WorkerCommand._execute(claimed.id, claimed.lease_token)
    job.refresh_from_db()
    return job


@pytest.mark.django_db(transaction=True)
def test_admin_upload_only_queues_then_worker_parses_all_articles(
    authenticated_client, admin_user
) -> None:
    response = upload_as_admin(authenticated_client, admin_user)
    assert response.data["processing_status"] == "pending"
    assert response.data["processing_job"]["status"] == "queued"
    assert response.data["processing_job"]["task_type"] == "knowledge_parse"
    assert response.data["source_verification"] == "pending"
    assert response.data["source_label"] == "待核验"
    assert response.data["clause_count"] == 0
    assert KnowledgeClauseRecord.objects.count() == 0

    parse_job = run_parse_job(document_id=response.data["id"])
    assert parse_job.status == ProcessingJob.Status.SUCCEEDED
    refreshed = authenticated_client(admin_user).get(
        f"/api/v1/knowledge/documents/{response.data['id']}/"
    )
    assert refreshed.status_code == 200
    assert refreshed.data["processing_status"] == "succeeded"
    assert refreshed.data["processing_job"]["status"] == "succeeded"
    assert refreshed.data["clause_count"] == 74
    assert KnowledgeClauseRecord.objects.count() == 74
    embedding_job = ProcessingJob.objects.get(
        task_type=ProcessingJob.TaskType.EMBEDDING,
        scope_type="knowledge_document",
        scope_id=response.data["id"],
    )
    assert embedding_job.inspection_id is None
    assert embedding_job.expected_scope_revision == refreshed.data["revision"]


@pytest.mark.django_db
def test_knowledge_upload_idempotency_does_not_duplicate_resource_or_job(
    authenticated_client, admin_user
) -> None:
    first = upload_as_admin(
        authenticated_client,
        admin_user,
        idempotency_key="same-knowledge-upload",
    )
    repeated = authenticated_client(admin_user).post(
        "/api/v1/knowledge/documents/",
        {
            "file": uploaded_fire_law(),
            "title": "中华人民共和国消防法（测试样本）",
            "category": "law",
        },
        format="multipart",
        HTTP_IDEMPOTENCY_KEY="same-knowledge-upload",
    )
    assert repeated.status_code == 200, repeated.data
    assert repeated.data["id"] == first.data["id"]
    assert repeated.data["processing_job"]["id"] == first.data["processing_job"]["id"]
    assert KnowledgeDocumentRecord.objects.count() == 1
    assert (
        ProcessingJob.objects.filter(task_type=ProcessingJob.TaskType.KNOWLEDGE_PARSE).count() == 1
    )

    conflict = authenticated_client(admin_user).post(
        "/api/v1/knowledge/documents/",
        {
            "file": uploaded_fire_law(),
            "title": "同一幂等键下的不同标题",
            "category": "law",
        },
        format="multipart",
        HTTP_IDEMPOTENCY_KEY="same-knowledge-upload",
    )
    assert conflict.status_code == 409
    assert conflict.data["error"]["code"] == "idempotency_conflict"


@pytest.mark.django_db(transaction=True)
def test_inspector_cannot_manage_knowledge_but_can_search_and_open_source(
    authenticated_client, admin_user, inspector_a
) -> None:
    created = upload_as_admin(authenticated_client, admin_user)
    run_parse_job(document_id=created.data["id"])
    client = authenticated_client(inspector_a)
    assert client.get("/api/v1/knowledge/documents/").status_code == 403
    assert (
        client.post(
            "/api/v1/knowledge/documents/",
            {"file": uploaded_fire_law(), "title": "重复", "category": "law"},
            format="multipart",
        ).status_code
        == 403
    )
    query = "自动消防系统操作人员必须持证上岗"
    response = client.get("/api/v1/knowledge/search/", {"q": query})
    assert response.status_code == 200
    assert response.data["found"] is True
    assert response.data["retrieval_mode"] == "lexical_fallback"
    assert response.data["embedding_status"] == "unavailable"
    assert response.data["embedding_fallback"] is True
    article = next(
        item for item in response.data["results"] if item["clause_identifier"] == "第二十一条"
    )
    assert article["clause_id"]
    assert article["issue"] == query
    assert article["document_id"] == created.data["id"]
    assert article["document_title"] == "中华人民共和国消防法（测试样本）"
    assert article["document_url"] == (
        f"/api/v1/knowledge/documents/{created.data['id']}/download/"
    )
    assert article["source_status"] == "待核验"
    assert article["applicability_reason"].strip()
    assert "人工判断适用性" in article["applicability_reason"]
    assert 1 <= article["page_start"] <= article["page_end"]
    detail = client.get(
        f"/api/v1/knowledge/documents/{created.data['id']}/clauses/{article['clause_id']}/"
    )
    assert detail.status_code == 200, detail.data
    assert detail.data["document_id"] == article["document_id"]
    assert detail.data["document_title"] == article["document_title"]
    assert detail.data["clause_id"] == article["clause_id"]
    assert detail.data["clause_identifier"] == article["clause_identifier"]
    assert detail.data["original_text"] == article["original_text"]
    assert detail.data["page_start"] == article["page_start"]
    assert detail.data["page_end"] == article["page_end"]
    assert detail.data["source_locator"] == {
        "page_start": article["page_start"],
        "page_end": article["page_end"],
    }
    assert detail.data["document_url"] == article["document_url"]
    assert detail.data["source_verification"] == "pending"
    assert detail.data["source_label"] == "待核验"
    assert detail.data["validity"] == "unknown"

    other_document = KnowledgeDocumentRecord.objects.create(
        title="其他来源文档",
        category=KnowledgeDocumentRecord.Category.OTHER,
        original_file="knowledge/other/original.pdf",
        mime_type="application/pdf",
        sha256="f" * 64,
        byte_size=1,
        processing_status=KnowledgeDocumentRecord.ProcessingStatus.SUCCEEDED,
        uploaded_by=admin_user,
    )
    mismatched = client.get(
        f"/api/v1/knowledge/documents/{other_document.id}/clauses/{article['clause_id']}/"
    )
    assert mismatched.status_code == 404

    download = client.get(f"/api/v1/knowledge/documents/{created.data['id']}/download/")
    assert download.status_code == 200
    downloaded_payload = b"".join(download.streaming_content)
    download.close()
    assert sha256(downloaded_payload).hexdigest() == PDF_SHA256
    assert created.data["sha256"] == PDF_SHA256


@pytest.mark.django_db(transaction=True)
def test_unrelated_query_is_explicit_not_found(authenticated_client, admin_user) -> None:
    created = upload_as_admin(authenticated_client, admin_user)
    run_parse_job(document_id=created.data["id"])
    response = authenticated_client(admin_user).get(
        "/api/v1/knowledge/search/", {"q": "量子传送门的星际消防许可规则"}
    )
    assert response.status_code == 200
    assert response.data["found"] is False
    assert response.data["results"] == []
    assert "未找到" in response.data["message"]


@pytest.mark.django_db
def test_internal_policy_cannot_be_uploaded_as_formal_law(authenticated_client, admin_user) -> None:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("第一条 内部值班要求")
    document.save(buffer)
    response = authenticated_client(admin_user).post(
        "/api/v1/knowledge/documents/",
        {
            "file": SimpleUploadedFile(
                "内部制度.docx",
                buffer.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            ),
            "title": "内部制度",
            "category": "law",
            "declared_internal": True,
        },
        format="multipart",
    )
    assert response.status_code == 400
    assert KnowledgeDocumentRecord.objects.count() == 0


@pytest.mark.django_db(transaction=True)
def test_parse_failure_preserves_original_and_exposes_retryable_failure(
    authenticated_client, admin_user, inspector_a
) -> None:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("没有条款标记的内部资料")
    document.save(buffer)
    response = authenticated_client(admin_user).post(
        "/api/v1/knowledge/documents/",
        {
            "file": SimpleUploadedFile(
                "资料.docx",
                buffer.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            ),
            "title": "待人工分段资料",
            "category": "other",
        },
        format="multipart",
        HTTP_IDEMPOTENCY_KEY="upload-unstructured-knowledge",
    )
    assert response.status_code == 202
    assert response.data["processing_status"] == "pending"
    record = KnowledgeDocumentRecord.objects.get(pk=response.data["id"])
    assert Path(record.original_file.path).is_file()

    job = run_parse_job(document_id=record.id)
    assert job.status == ProcessingJob.Status.FAILED
    failed = authenticated_client(admin_user).get(f"/api/v1/knowledge/documents/{record.id}/")
    assert failed.data["processing_status"] == "failed"
    assert failed.data["safe_error_message"]
    assert failed.data["processing_job"]["safe_error_message"]

    retried = authenticated_client(admin_user).post(
        f"/api/v1/knowledge/documents/{record.id}/retry/",
        {},
        format="json",
        HTTP_IDEMPOTENCY_KEY="retry-unstructured-knowledge",
    )
    assert retried.status_code == 202
    assert retried.data["processing_status"] == "pending"
    assert retried.data["processing_job"]["id"] == str(job.id)
    assert retried.data["processing_job"]["status"] == "queued"
    rerun = run_parse_job(document_id=record.id)
    assert rerun.id == job.id
    assert rerun.status == ProcessingJob.Status.FAILED
    replayed_retry = authenticated_client(admin_user).post(
        f"/api/v1/knowledge/documents/{record.id}/retry/",
        {},
        format="json",
        HTTP_IDEMPOTENCY_KEY="retry-unstructured-knowledge",
    )
    assert replayed_retry.status_code == 200
    assert replayed_retry.data["processing_status"] == "failed"
    assert replayed_retry.data["processing_job"]["status"] == "failed"
    assert (
        ProcessingJob.objects.filter(
            task_type=ProcessingJob.TaskType.KNOWLEDGE_PARSE,
            scope_id=record.id,
        ).count()
        == 1
    )
    assert Path(record.original_file.path).is_file()
    forbidden = authenticated_client(inspector_a).post(
        f"/api/v1/knowledge/documents/{record.id}/retry/",
        {},
        format="json",
        HTTP_IDEMPOTENCY_KEY="forbidden-retry",
    )
    assert forbidden.status_code == 403
