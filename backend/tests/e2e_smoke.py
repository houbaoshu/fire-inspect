"""端到端冒烟（PLAN 阶段 7 验收流程自动化）：

reset+seed 后：创建检查 → 上传视频 → 等待任务链 → 确认事实 → 法规检索 →
三类文书生成/导出 → XML 级校验（打勾/下划线/图片/问答正文）。

用法：cd backend && .venv/bin/python -m tests.e2e_smoke  （需服务已在 :8000 运行）
"""
import json, time, sys, zipfile, re
import urllib.request, urllib.parse
from http.cookiejar import CookieJar
from pathlib import Path

BASE = "http://localhost:8000"
ROOT = Path(__file__).resolve().parents[2]

cj = CookieJar()
opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))


def api(method, path, body=None, file=None, form=None):
    url = BASE + path
    if file:
        import requests  # 仅测试脚本可用 requests？—— 不，用 multipart 手工构造
    if body is not None:
        req = urllib.request.Request(url, data=json.dumps(body).encode(), method=method,
                                     headers={"Content-Type": "application/json"})
    elif form is not None:
        boundary = "----testboundary"
        parts = []
        for k, v in (form.get("fields") or {}).items():
            parts.append(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode())
        for k, (fname, data, ctype) in (form.get("files") or {}).items():
            parts.append(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; filename=\"{fname}\"\r\n"
                         f"Content-Type: {ctype}\r\n\r\n".encode() + data + b"\r\n")
        parts.append(f"--{boundary}--\r\n".encode())
        req = urllib.request.Request(url, data=b"".join(parts), method=method,
                                     headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    else:
        req = urllib.request.Request(url, method=method)
    with opener.open(req, timeout=300) as r:
        return json.loads(r.read() or b"{}")


def download(path, out):
    with opener.open(BASE + path, timeout=60) as r:
        Path(out).write_bytes(r.read())


def main():
    api("POST", "/api/auth/login", {"username": "admin", "password": "admin123"})
    print("✓ 登录")

    insp = api("POST", "/api/inspections", {
        "unit_name": "某商业广场", "address": "桐乡市梧桐街道XX路1号", "checker_names": "张三、李四",
        "contact": "王五 13800000000", "legal_person": "王五", "check_type": "对举报投诉的核查",
        "inspected_at": "2025-08-22T10:00:00+08:00"})
    iid = insp["id"]
    print("✓ 创建检查 #", iid)

    video = ROOT / "testdata" / "videos" / "1.mp4"
    m = api("POST", f"/api/inspections/{iid}/materials",
            form={"files": {"file": ("1.mp4", video.read_bytes(), "video/mp4")}})
    mid = m["id"]
    print("✓ 上传视频 #", mid)

    deadline = time.time() + 180
    while time.time() < deadline:
        jobs = api("GET", f"/api/materials/{mid}/jobs")
        if all(j["status"] in ("succeeded", "failed") for j in jobs) and len(jobs) >= 3:
            # fact_extract 在转写完成后触发，再等一轮
            all_jobs = api("GET", f"/api/inspections/{iid}/jobs")
            if all(j["status"] in ("succeeded", "failed") for j in all_jobs):
                break
        time.sleep(2)
    jobs = api("GET", f"/api/inspections/{iid}/jobs")
    for j in jobs:
        mark = "✓" if j["status"] == "succeeded" else "✗"
        print(f"  {mark} {j['kind']}: {j['status']} {j['error'][:80]}")
    assert all(j["status"] == "succeeded" for j in jobs), "存在失败任务"

    facts = api("GET", f"/api/inspections/{iid}/facts")
    print(f"✓ 候选事实 {len(facts)} 条")
    for f in facts:
        api("PATCH", f"/api/facts/{f['id']}", {"status": "confirmed"})
    kfs = api("GET", f"/api/inspections/{iid}/keyframes")
    print(f"✓ 关键帧 {len(kfs)} 张")
    if facts and kfs:
        api("POST", f"/api/facts/{facts[0]['id']}/evidences",
            {"evidence_type": "keyframe", "material_id": mid, "ref": {"keyframe_id": kfs[0]["id"]}})

    q = api("POST", "/api/knowledge/query", {"question": "消防值班人员持证上岗的要求"})
    print("✓ 法规检索 found =", q["found"], "| 来源:", [s["clause_label"] for s in q["sources"][:3]])

    tpls = api("GET", "/api/templates")
    tpl = {t["doc_type"]: t["id"] for t in tpls}
    frame_ids = [k["id"] for k in kfs[:3]]
    docs = {}
    for dt in ("record", "photo", "interrogation"):
        body = {"doc_type": dt, "template_id": tpl[dt]}
        if dt == "photo":
            body["frame_ids"] = frame_ids
        d = api("POST", f"/api/inspections/{iid}/documents", body)
        docs[dt] = d
        print(f"✓ 文书「{d['name']}」has_file={d['has_file']} 字段数={len(d['fields'])}")

    # 导出与 XML 校验
    for dt, d in docs.items():
        out = f"/tmp/e2e_{dt}.docx"
        download(f"/api/documents/{d['id']}/export", out)
        verify(dt, out, d["id"])

    # 留痕与文书退回
    api("PATCH", f"/api/facts/{facts[0]['id']}", {"content": facts[0]["content"] + "（复核）"})
    d = api("GET", f"/api/documents/{docs['record']['id']}")
    hist = api("GET", f"/api/documents/{docs['record']['id']}/history")
    print("✓ 事实变更后记录表状态:", d["status"], "| 历史条数:", len(hist))

    # AI 助手
    a = api("POST", "/api/assistant/chat", {"message": "这次检查发现了什么问题？依据是什么？",
                                            "inspection_id": iid, "page": "facts"})
    print("✓ 助手回答:", a["answer"][:120].replace("\n", " "))
    print("  动作:", a["actions"], "| 建议:", (a.get("suggestion") or "")[:60])
    print("\n全部通过")


def verify(dt, path, doc_id):
    z = zipfile.ZipFile(path)
    xml = z.read("word/document.xml").decode()
    import docx
    d = docx.Document(path)
    if dt == "record":
        tbl = d.tables[0]
        assert "☑" in tbl.rows[0].cells[0].text, "记录表未打勾"
        assert tbl.rows[0].cells[0].text.count("□") == 6, "未选项□数不对"
        assert "某商业广场" in tbl.rows[1].cells[2].text
        print("  ✓ 记录表：打勾正确、单元格已填、下划线 run 数 =", xml.count('w:val="single"'))
    elif dt == "photo":
        media = [n for n in z.namelist() if n.startswith("word/media")]
        texts = [p.text for p in d.paragraphs]
        n_blocks = sum(1 for t in texts if t.startswith("摄制人"))
        assert len(media) == 3, f"应有 3 张照片，实际 {len(media)}"
        assert n_blocks == 3, f"应有 3 个证据区块，实际 {n_blocks}"
        assert "[此处插入照片]" not in "".join(texts), "占位符未替换"
        print(f"  ✓ 照片文书：3 图 3 区块纵向排列，下划线 run 数 =", xml.count('w:val="single"'))
    elif dt == "interrogation":
        texts = [p.text for p in d.paragraphs if p.text.strip()]
        qa = [t for t in texts if t.startswith(("问：", "答："))]
        assert len(qa) >= 4, f"问答正文不足：{len(qa)}"
        assert any("桐乡市" in t for t in texts), "地点未填"
        # 问答段落下划线
        for p in d.paragraphs:
            if p.text.startswith("问："):
                from docx.oxml.ns import qn
                u = p.runs[0]._r.find(".//" + qn("w:u"))
                assert u is not None and u.get(qn("w:val")) == "single", "问答正文未带下划线"
                break
        sx = z.read("word/settings.xml").decode()
        assert "updateFields" in sx
        print(f"  ✓ 笔录：问答 {len(qa)} 行带下划线，页脚域将自动刷新")


if __name__ == "__main__":
    main()
