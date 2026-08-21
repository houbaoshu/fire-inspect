import base64
import io
import json
import shutil
import subprocess
import zipfile

from app.core.database import SessionLocal
from app.models import Fact
from conftest import login
from docx import Document


def create_inspection(client, title="某单位消防监督检查"):
    response = client.post(
        "/api/inspections",
        json={
            "title": title,
            "extra_fields": {
                "organization": "示例消防安全重点单位",
                "address": "示例路 1 号",
                "inspection_date": "2026-08-21 10:00",
                "inspectors": "张检查员、李检查员",
            },
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_user_session_and_inspection_isolation(client):
    login(client)
    for username in ("inspector-a", "inspector-b"):
        response = client.post(
            "/api/users",
            json={"username": username, "password": "safe-password", "role": "inspector"},
        )
        assert response.status_code == 201
    client.post("/api/auth/logout")

    login(client, "inspector-a", "safe-password")
    inspection = create_inspection(client)
    client.post("/api/auth/logout")

    login(client, "inspector-b", "safe-password")
    response = client.get(f"/api/inspections/{inspection['id']}")
    assert response.status_code == 404


def test_ai_fact_cannot_be_confirmed_without_evidence(client):
    login(client)
    inspection = create_inspection(client)
    with SessionLocal() as db:
        fact = Fact(
            inspection_id=inspection["id"],
            kind="observation",
            summary="没有证据定位的 AI 事实",
            source="ai",
            status="candidate",
        )
        db.add(fact)
        db.commit()
        fact_id = fact.id

    response = client.post(f"/api/facts/{fact_id}/confirm")
    assert response.status_code == 422
    assert "缺少证据" in response.text


def test_knowledge_query_returns_resolvable_citation(client):
    login(client)
    content = (
        "消防控制室应当实行每日二十四小时专人值班制度，每班不应少于二人。\n"
        "值班人员应持有相应职业资格并做好值班记录。"
    ).encode()
    response = client.post(
        "/api/knowledge-documents",
        data={"title": "消防控制室值班参考规定", "material_type": "formal"},
        files={"file": ("control-room.txt", content, "text/plain")},
    )
    assert response.status_code == 201, response.text
    document_id = response.json()["document"]["id"]

    response = client.post(
        "/api/knowledge/query",
        json={"question": "消防控制室无人值班可能涉及什么规定？"},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["characterization"] != "未找到足够依据"
    assert result["citations"]
    uploaded_citation = next(
        item for item in result["citations"] if item["document_id"] == document_id
    )
    assert "二十四小时专人值班" in uploaded_citation["excerpt"]
    assert client.get(f"/api/knowledge-documents/{document_id}/content").status_code == 200


def test_confirmed_fact_to_editable_and_exportable_word_document(client):
    login(client)
    inspection = create_inspection(client)
    response = client.post(
        f"/api/inspections/{inspection['id']}/facts",
        json={
            "kind": "observation",
            "summary": "检查发现消防控制室无人值班。",
            "content": {"value": "检查发现消防控制室无人值班。"},
            "evidence_ids": [],
        },
    )
    assert response.status_code == 201
    fact = response.json()
    assert fact["source"] == "user"
    assert fact["status"] == "confirmed"

    templates = client.get("/api/templates").json()
    assert {item["document_type"] for item in templates} == {
        "inspection_record",
        "photo_record",
        "inquiry_record",
    }
    template = next(item for item in templates if item["document_type"] == "inspection_record")
    response = client.post(
        f"/api/inspections/{inspection['id']}/documents",
        json={"template_id": template["id"]},
    )
    assert response.status_code == 201, response.text
    draft = response.json()
    assert "消防控制室无人值班" in draft["fields_json"]["facts"]["value"]
    assert draft["fields_json"]["facts"]["fact_ids"] == [fact["id"]]

    edited_value = "经人工核对：检查时消防控制室无人值班。"
    response = client.patch(
        f"/api/documents/{draft['id']}/fields/facts",
        json={"value": edited_value},
    )
    assert response.status_code == 200
    assert response.json()["fields_json"]["facts"]["edited"] is True

    assert client.post(f"/api/documents/{draft['id']}/confirm").status_code == 200
    response = client.post(f"/api/documents/{draft['id']}/export")
    assert response.status_code == 200, response.text
    download = client.get(response.json()["download_url"])
    assert download.status_code == 200
    assert download.content.startswith(b"PK")
    with zipfile.ZipFile(io.BytesIO(download.content)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert edited_value in xml
    assert "{{" not in xml


def test_fact_change_marks_generated_field_stale_without_overwriting(client):
    login(client)
    inspection = create_inspection(client)
    fact = client.post(
        f"/api/inspections/{inspection['id']}/facts",
        json={"kind": "observation", "summary": "原始事实", "content": {}, "evidence_ids": []},
    ).json()
    template = next(
        item
        for item in client.get("/api/templates").json()
        if item["document_type"] == "inspection_record"
    )
    draft = client.post(
        f"/api/inspections/{inspection['id']}/documents",
        json={"template_id": template["id"]},
    ).json()
    manual_value = "人工已经修改的事实描述"
    client.patch(f"/api/documents/{draft['id']}/fields/facts", json={"value": manual_value})
    client.patch(f"/api/facts/{fact['id']}", json={"summary": "修正后的事实"})
    current = client.get(f"/api/inspections/{inspection['id']}/documents").json()[0]
    assert current["fields_json"]["facts"]["value"] == manual_value
    assert current["fields_json"]["facts"]["stale"] is True
    assert client.post(f"/api/documents/{draft['id']}/confirm").status_code == 422


def test_new_fact_invalidates_confirmed_document_without_overwriting(client):
    login(client)
    inspection = create_inspection(client)
    template = next(
        item
        for item in client.get("/api/templates").json()
        if item["document_type"] == "inspection_record"
    )
    draft = client.post(
        f"/api/inspections/{inspection['id']}/documents",
        json={"template_id": template["id"]},
    ).json()
    assert client.post(f"/api/documents/{draft['id']}/confirm").status_code == 200

    response = client.post(
        f"/api/inspections/{inspection['id']}/facts",
        json={"kind": "observation", "summary": "新确认事实", "content": {}, "evidence_ids": []},
    )
    assert response.status_code == 201
    current = client.get(f"/api/inspections/{inspection['id']}/documents").json()[0]
    assert current["status"] == "draft"
    assert current["fields_json"]["facts"]["stale"] is True
    assert current["fields_json"]["facts"]["value"] == ""
    assert client.post(f"/api/documents/{draft['id']}/export").status_code == 422


def test_short_video_upload_extracts_traceable_evidence_without_inventing_facts(client, tmp_path):
    if not shutil.which("ffmpeg"):
        return
    login(client)
    inspection = create_inspection(client)
    video = tmp_path / "short.mp4"
    subprocess.run(
        [
            "ffmpeg",
            "-y",
            "-f",
            "lavfi",
            "-i",
            "color=c=red:s=640x360:d=1",
            "-f",
            "lavfi",
            "-i",
            "anullsrc=r=16000:cl=mono",
            "-shortest",
            "-c:v",
            "libx264",
            "-c:a",
            "aac",
            str(video),
        ],
        check=True,
        capture_output=True,
    )
    with video.open("rb") as fileobj:
        response = client.post(
            f"/api/inspections/{inspection['id']}/assets",
            files={"file": (video.name, fileobj, "video/mp4")},
        )
    assert response.status_code == 201, response.text
    assert response.json()["asset"]["processing_status"] == "completed", response.json()[
        "asset"
    ].get("processing_error")
    assets = client.get(f"/api/inspections/{inspection['id']}/assets").json()
    assert "未生成候选事实" in assets[0]["metadata_json"]["warning"]
    evidence = client.get(f"/api/inspections/{inspection['id']}/evidence").json()
    assert any(item["kind"] == "keyframe" and item["start_ms"] == 0 for item in evidence)
    assert client.get(f"/api/inspections/{inspection['id']}/facts").json() == []


def test_photo_and_inquiry_documents_export_from_shared_facts(client):
    login(client)
    inspection = create_inspection(client)
    client.post(
        f"/api/inspections/{inspection['id']}/facts",
        json={
            "kind": "observation",
            "summary": "疏散通道堆放物品。",
            "content": {},
            "evidence_ids": [],
        },
    )
    qa_text = "问：现场由谁负责消防安全管理？\n答：由单位安全负责人负责。"
    client.post(
        f"/api/inspections/{inspection['id']}/facts",
        json={"kind": "qa", "summary": qa_text, "content": {}, "evidence_ids": []},
    )
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Z2S8AAAAASUVORK5CYII="
    )
    response = client.post(
        f"/api/inspections/{inspection['id']}/assets",
        files={"file": ("evidence.png", png, "image/png")},
    )
    assert response.status_code == 201
    photo = next(
        item
        for item in client.get(f"/api/inspections/{inspection['id']}/evidence").json()
        if item["kind"] == "photo"
    )
    templates = {item["document_type"]: item for item in client.get("/api/templates").json()}
    for document_type in ("photo_record", "inquiry_record"):
        draft = client.post(
            f"/api/inspections/{inspection['id']}/documents",
            json={"template_id": templates[document_type]["id"]},
        ).json()
        if document_type == "photo_record":
            client.patch(
                f"/api/documents/{draft['id']}/fields/photo",
                json={"value": photo["id"]},
            )
        else:
            assert qa_text in draft["fields_json"]["questions_answers"]["value"]
        assert client.post(f"/api/documents/{draft['id']}/confirm").status_code == 200
        result = client.post(f"/api/documents/{draft['id']}/export")
        assert result.status_code == 200, result.text
        content = client.get(result.json()["download_url"]).content
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
            assert "{{" not in xml
            if document_type == "photo_record":
                assert any(name.startswith("word/media/") for name in archive.namelist())
            else:
                assert "单位安全负责人" in xml


def test_uploaded_template_fields_can_be_configured_without_changing_codes(client):
    login(client)
    document = Document()
    document.add_paragraph("自定义字段：{{ custom_field }}")
    output = io.BytesIO()
    document.save(output)
    response = client.post(
        "/api/templates",
        data={"name": "自定义检查记录", "document_type": "inspection_record"},
        files={
            "file": (
                "custom.docx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    template = response.json()
    response = client.patch(
        f"/api/templates/{template['id']}/fields",
        json={
            "fields": [
                {
                    "code": "custom_field",
                    "label": "自定义事实",
                    "type": "multiline",
                    "description": "只根据已确认事实生成",
                }
            ]
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["fields_json"][0]["label"] == "自定义事实"
    response = client.patch(
        f"/api/templates/{template['id']}/fields",
        json={
            "fields": [
                {
                    "code": "renamed",
                    "label": "错误字段",
                    "type": "text",
                    "description": "",
                }
            ]
        },
    )
    assert response.status_code == 422


def test_photo_list_field_exports_multiple_images(client):
    login(client)
    inspection = create_inspection(client)
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Z2S8AAAAASUVORK5CYII="
    )
    photo_ids = []
    for filename in ("first.png", "second.png"):
        response = client.post(
            f"/api/inspections/{inspection['id']}/assets",
            files={"file": (filename, png, "image/png")},
        )
        assert response.status_code == 201
        photo_ids.append(
            next(
                item["id"]
                for item in client.get(f"/api/inspections/{inspection['id']}/evidence").json()
                if item["kind"] == "photo" and item["id"] not in photo_ids
            )
        )

    document = Document()
    document.add_paragraph("现场照片：{{ photos }}")
    output = io.BytesIO()
    document.save(output)
    template = client.post(
        "/api/templates",
        data={
            "name": "多图照片记录",
            "document_type": "photo_record",
            "fields_json": json.dumps(
                [
                    {
                        "code": "photos",
                        "label": "现场照片",
                        "type": "photo_list",
                        "description": "选择多张现场照片",
                    }
                ]
            ),
        },
        files={
            "file": (
                "photo-list.docx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()
    draft = client.post(
        f"/api/inspections/{inspection['id']}/documents",
        json={"template_id": template["id"]},
    ).json()
    assert draft["fields_json"]["photos"]["value"] == []
    response = client.patch(
        f"/api/documents/{draft['id']}/fields/photos",
        json={"value": photo_ids},
    )
    assert response.status_code == 200, response.text
    assert response.json()["fields_json"]["photos"]["evidence_ids"] == photo_ids
    assert client.post(f"/api/documents/{draft['id']}/confirm").status_code == 200
    exported = client.post(f"/api/documents/{draft['id']}/export").json()
    content = client.get(exported["download_url"]).content
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        assert xml.count("<w:drawing>") == 2
