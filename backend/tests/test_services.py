import io

import pytest
from app.core.config import Settings
from app.models import Asset, Evidence
from app.services.ai import AIProvider
from app.services.documents import scan_template_fields
from app.services.knowledge import extract_document, split_chunks
from app.services.media import _persist_fact_candidates
from docx import Document


def test_template_placeholder_scan():
    document = Document()
    document.add_paragraph("单位：{{ organization }}")
    document.add_paragraph("事实：{{ facts }}")
    output = io.BytesIO()
    document.save(output)
    assert scan_template_fields(output.getvalue()) == ["facts", "organization"]


def test_template_placeholder_scan_in_table_and_header():
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "单位：{{ organization }}"
    document.sections[0].header.paragraphs[0].text = "检查：{{ inspection_title }}"
    output = io.BytesIO()
    document.save(output)
    assert scan_template_fields(output.getvalue()) == ["inspection_title", "organization"]


def test_template_rejects_non_field_jinja_expressions():
    document = Document()
    document.add_paragraph("{{ organization }}")
    document.add_paragraph("{{ unsafe.attribute }}")
    output = io.BytesIO()
    document.save(output)
    with pytest.raises(ValueError, match="只能使用"):
        scan_template_fields(output.getvalue())


def test_text_knowledge_extraction_and_chunking():
    sections = extract_document("第一条 消防安全要求\n第二条 值班要求".encode(), "rules.txt")
    chunks = split_chunks(sections, max_chars=20)
    assert chunks
    assert all(locator == "全文" for locator, _ in chunks)
    assert "消防安全要求" in "".join(content for _, content in chunks)


class FakeSession:
    def __init__(self):
        self.added = []

    def add(self, item):
        self.added.append(item)

    def flush(self):
        return None


@pytest.mark.parametrize(
    "extracted",
    [
        {
            "facts": [
                {
                    "kind": "observation",
                    "summary": "疏散通道被占用",
                    "content": {},
                    "confidence": 0.9,
                    "evidence_ids": ["evidence-1"],
                }
            ]
        },
        [
            {
                "kind": "observation",
                "summary": "疏散通道被占用",
                "content": {},
                "confidence": 0.9,
                "evidence_ids": ["evidence-1"],
            }
        ],
    ],
)
def test_fact_candidates_accept_object_or_top_level_array(extracted):
    db = FakeSession()
    asset = Asset(inspection_id="inspection-1")
    evidence = [Evidence(id="evidence-1", inspection_id="inspection-1")]

    assert _persist_fact_candidates(db, asset, evidence, extracted) == 1
    assert any(item.summary == "疏散通道被占用" for item in db.added if hasattr(item, "summary"))


def test_ai_provider_routes_capabilities_to_separate_base_urls(monkeypatch, tmp_path):
    calls = []

    class Response:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self.payload

    def fake_post(url, **kwargs):
        calls.append((url, kwargs))
        if url.endswith("/embeddings"):
            return Response({"data": [{"embedding": [0.0] * 256}]})
        if url.endswith("/audio/transcriptions"):
            return Response({"segments": [{"start": 0, "end": 1, "text": "测试"}]})
        return Response({"choices": [{"message": {"content": '{"facts":[]}'}}]})

    monkeypatch.setattr("app.services.ai.httpx.post", fake_post)
    provider = AIProvider()
    provider.settings = Settings(
        ai_base_url="https://legacy.example/v1",
        ai_chat_base_url="https://chat.example/v1/",
        ai_transcription_base_url="https://speech.example/v1",
        ai_embedding_base_url="https://embedding.example/v1",
        ai_api_key="test-key",
        ai_chat_model="chat-model",
        ai_transcription_model="speech-model",
        ai_embedding_model="embedding-model",
        ai_embedding_dimensions=256,
    )
    audio = tmp_path / "audio.wav"
    audio.write_bytes(b"test")

    provider.chat_json(system="system", user="user")
    provider.transcribe(audio)
    provider.embeddings(["text"])

    assert [url for url, _ in calls] == [
        "https://chat.example/v1/chat/completions",
        "https://speech.example/v1/audio/transcriptions",
        "https://embedding.example/v1/embeddings",
    ]
    assert calls[2][1]["json"]["dimensions"] == 256


def test_ai_provider_uses_aliyun_native_api_for_qwen_audio(monkeypatch, tmp_path):
    calls = []

    class Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {
                "output": {
                    "sentence": {"begin_time": 100, "end_time": 3800, "text": "测试"},
                    "text": "完整测试文本",
                },
                "usage": {"duration": 4},
            }

    def fake_post(url, **kwargs):
        calls.append((url, kwargs))
        return Response()

    monkeypatch.setattr("app.services.ai.httpx.post", fake_post)
    provider = AIProvider()
    provider.settings = Settings(
        ai_transcription_base_url=(
            "https://workspace.cn-beijing.maas.aliyuncs.com/compatible-mode/v1"
        ),
        ai_api_key="test-key",
        ai_transcription_model="qwen-audio-3.0-asr-flash",
    )
    audio = tmp_path / "audio.wav"
    audio.write_bytes(b"test")

    assert provider.transcribe(audio) == [
        {"start_ms": 0, "end_ms": 4000, "text": "完整测试文本"}
    ]
    assert calls[0][0] == (
        "https://workspace.cn-beijing.maas.aliyuncs.com/"
        "api/v1/services/aigc/multimodal-generation/generation"
    )
    request = calls[0][1]
    assert request["headers"]["X-DashScope-SSE"] == "disable"
    assert request["json"]["model"] == "qwen-audio-3.0-asr-flash"
    assert request["json"]["parameters"] == {"format": "wav", "sample_rate": "16000"}
    assert (
        request["json"]["input"]["messages"][0]["content"][0]["input_audio"]["data"]
        == "data:audio/wav;base64,dGVzdA=="
    )


def test_ai_provider_uses_chat_completions_for_qwen3_asr(monkeypatch, tmp_path):
    calls = []

    class Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {
                "choices": [{"message": {"content": "转写文本"}}],
                "usage": {"seconds": 2.5},
            }

    def fake_post(url, **kwargs):
        calls.append((url, kwargs))
        return Response()

    monkeypatch.setattr("app.services.ai.httpx.post", fake_post)
    provider = AIProvider()
    provider.settings = Settings(
        ai_transcription_base_url="https://workspace.cn-beijing.maas.aliyuncs.com/api/v1",
        ai_api_key="test-key",
        ai_transcription_model="qwen3-asr-flash",
    )
    audio = tmp_path / "audio.wav"
    audio.write_bytes(b"test")

    assert provider.transcribe(audio) == [
        {"start_ms": 0, "end_ms": 2500, "text": "转写文本"}
    ]
    assert calls[0][0] == (
        "https://workspace.cn-beijing.maas.aliyuncs.com/compatible-mode/v1/chat/completions"
    )
    assert calls[0][1]["json"]["messages"][0]["content"][0]["type"] == "input_audio"


def test_ai_provider_rejects_async_filetrans_model(tmp_path):
    provider = AIProvider()
    provider.settings = Settings(
        ai_transcription_base_url="https://workspace.cn-beijing.maas.aliyuncs.com/api/v1",
        ai_api_key="test-key",
        ai_transcription_model="qwen3-asr-flash-filetrans",
    )
    audio = tmp_path / "audio.wav"
    audio.write_bytes(b"test")

    with pytest.raises(ValueError, match="仅支持异步文件转写"):
        provider.transcribe(audio)


def test_ai_capability_urls_fall_back_to_legacy_base_url():
    settings = Settings(
        ai_base_url="https://legacy.example/v1",
        ai_chat_base_url=None,
        ai_transcription_base_url=None,
        ai_embedding_base_url=None,
    )
    assert settings.chat_base_url == "https://legacy.example/v1"
    assert settings.transcription_base_url == "https://legacy.example/v1"
    assert settings.embedding_base_url == "https://legacy.example/v1"
