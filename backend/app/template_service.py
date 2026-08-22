from __future__ import annotations

import hashlib
import io
import re
import subprocess
import tempfile
import uuid
from collections.abc import Iterable
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import (
    DocumentField,
    DocumentType,
    ProcessingStatus,
    Template,
    TemplateRegion,
    TemplateVersion,
)
from .storage import get_storage

PLACEHOLDER = re.compile(r"\{\{\s*([\w.:-]+)\s*\}\}")
BLANK = re.compile(r"_{3,}|＿{3,}")
LABEL = re.compile(r"(?P<label>[\u4e00-\u9fffA-Za-z0-9（）()·]{1,30})[：:]")
PHOTO_PLACEHOLDER = re.compile(
    r"[\[【（(]?\s*此处\s*(?:插入|粘贴|放置)?\s*(?:照片|图片|图像)\s*[\]】）)]?"
)


@dataclass(frozen=True)
class ParagraphRef:
    path: str
    paragraph: Paragraph
    context_label: str = ""


def normalize_word(data: bytes, filename: str) -> bytes:
    if filename.lower().endswith(".docx"):
        Document(io.BytesIO(data))
        return data
    if not filename.lower().endswith(".doc"):
        raise ValueError("模板仅支持 DOCX 或 DOC")
    with tempfile.TemporaryDirectory() as temp_dir:
        source = Path(temp_dir) / "template.doc"
        source.write_bytes(data)
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", temp_dir, source],
            check=True,
            capture_output=True,
        )
        converted = Path(temp_dir) / "template.docx"
        if not converted.exists():
            raise ValueError("旧版 DOC 转换失败")
        return converted.read_bytes()


def _paragraph_refs(document: DocumentObject) -> Iterable[ParagraphRef]:
    for index, paragraph in enumerate(document.paragraphs):
        yield ParagraphRef(f"body:p:{index}", paragraph)
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            cells = row.cells
            for cell_index, cell in enumerate(cells):
                previous = cells[cell_index - 1].text.strip() if cell_index else ""
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield ParagraphRef(
                        f"body:table:{table_index}:row:{row_index}:cell:{cell_index}:p:{paragraph_index}",
                        paragraph,
                        previous,
                    )
    for section_index, section in enumerate(document.sections):
        for area_name, area in (("header", section.header), ("footer", section.footer)):
            for paragraph_index, paragraph in enumerate(area.paragraphs):
                yield ParagraphRef(f"{area_name}:{section_index}:p:{paragraph_index}", paragraph)


def _suggest_source(
    label: str,
    key: str,
    document_type: DocumentType | None = None,
    context: str = "",
) -> tuple[str, str]:
    value = f"{label} {key}".lower()
    normalized_label = re.sub(r"\s+", "", label).strip("：:")
    if document_type == DocumentType.INSPECTION_RECORD:
        inspection_sources = {
            "单位（场所）名称": ("inspection.unit_name", "text"),
            "单位(场所)名称": ("inspection.unit_name", "text"),
            "法定代表/主要负责人": ("custom.legal_representative", "text"),
            "地址": ("inspection.address", "text"),
            "联系人及联系电话": ("custom.contact_and_phone", "text"),
            "检查内容和情况": ("facts.confirmed_text", "multiline"),
            "备注": ("custom.remarks", "multiline"),
            "检查时间": ("inspection.inspected_at", "date"),
            "消防监督检查员（签名）": ("custom.inspector_signature", "text"),
            "消防监督检查员(签名)": ("custom.inspector_signature", "text"),
            "被检查单位随同检查人员（签名）": (
                "custom.accompanying_person_signature",
                "text",
            ),
            "被检查单位随同检查人员(签名)": (
                "custom.accompanying_person_signature",
                "text",
            ),
        }
        if normalized_label in inspection_sources:
            return inspection_sources[normalized_label]
    if document_type == DocumentType.PHOTO_DOCUMENT:
        photo_sources = {
            "反映内容": ("facts.confirmed_text", "multiline"),
            "拍摄地点": ("inspection.address", "text"),
            "摄制人": ("custom.photographer", "text"),
            "摄影人": ("custom.photographer", "text"),
            "拍摄时间": ("inspection.inspected_at", "date"),
            "说明": ("custom.photo_description", "multiline"),
            "确认人": ("custom.photo_confirmer", "text"),
        }
        if normalized_label in photo_sources:
            return photo_sources[normalized_label]
    if document_type == DocumentType.INTERVIEW_TRANSCRIPT:
        interview_sources = {
            "时间": ("custom.interview_time_range", "text"),
            "地点": ("inspection.address", "text"),
            "询问人（签名）": ("custom.questioner_name", "text"),
            "询问人(签名)": ("custom.questioner_name", "text"),
            "记录人（签名）": ("custom.recorder_name", "text"),
            "记录人(签名)": ("custom.recorder_name", "text"),
            "被询问人姓名": ("custom.interviewee_name", "text"),
            "性别": ("custom.interviewee_gender", "text"),
            "年龄": ("custom.interviewee_age", "text"),
            "出生日期": ("custom.interviewee_birth_date", "date"),
            "证件类型": ("custom.interviewee_id_type", "text"),
            "证件号码": ("custom.interviewee_id_number", "text"),
            "户籍所在地": ("custom.interviewee_registered_address", "text"),
            "现住址": ("custom.interviewee_address", "text"),
            "联系方式": ("custom.interviewee_contact", "text"),
            "问答": ("transcript.confirmed_text", "multiline"),
            "询问内容": ("transcript.confirmed_text", "multiline"),
        }
        if normalized_label == "工作单位":
            if "询问人" in context:
                return "custom.questioner_work_unit", "text"
            if "记录人" in context:
                return "custom.recorder_work_unit", "text"
            return "custom.interviewee_work_unit", "text"
        if normalized_label in interview_sources:
            return interview_sources[normalized_label]
    if "单位" in value or "unit" in value:
        return "inspection.unit_name", "text"
    if "地址" in value or "address" in value:
        return "inspection.address", "text"
    if "时间" in value or "日期" in value or "date" in value or "time" in value:
        return "inspection.inspected_at", "date"
    if "照片" in value or "photo" in value or "image" in value:
        return "document.photo_entries", "image"
    if "询问" in value or "问答" in value or "transcript" in value:
        return "transcript.confirmed_text", "multiline"
    if "情况" in value or "事实" in value or "结果" in value or "fact" in value:
        return "facts.confirmed_text", "multiline"
    safe_key = re.sub(r"[^a-zA-Z0-9_.-]+", "_", key).strip("_").lower()
    if not safe_key:
        safe_key = hashlib.sha1(label.encode()).hexdigest()[:12]
    return f"custom.{safe_key}", "text"


def _region_values(
    *,
    path: str,
    label: str,
    source: str,
    field_type: str,
    confidence: float,
    region_data: dict[str, Any],
    ordinal: int,
    repeatable: bool = False,
) -> dict[str, Any]:
    return {
        "path": path,
        "label": label,
        "field_key": source.removeprefix("custom."),
        "source_path": source,
        "field_type": field_type,
        "confidence": confidence,
        "region_data": region_data,
        "ordinal": ordinal,
        "repeatable": repeatable,
    }


def _label_regions(
    ref: ParagraphRef,
    document_type: DocumentType,
    ordinal: int,
) -> list[dict[str, Any]]:
    text = ref.paragraph.text
    matches = list(LABEL.finditer(text))
    regions: list[dict[str, Any]] = []
    for match_index, match in enumerate(matches):
        end = matches[match_index + 1].start() if match_index + 1 < len(matches) else len(text)
        gap = text[match.end() : end]
        label = match.group("label").strip()
        # Fixed prose after a label is not an empty filling area. Interview forms
        # commonly print only "至" in the otherwise empty time-range field.
        fixed_time_area = (
            document_type == DocumentType.INSPECTION_RECORD
            and label == "检查时间"
            and all(marker in gap for marker in ("年", "月", "日"))
        )
        interview_time_area = (
            document_type == DocumentType.INTERVIEW_TRANSCRIPT
            and label == "时间"
            and gap.strip() == "至"
        )
        if gap.strip() and not (fixed_time_area or interview_time_area):
            continue
        source, field_type = _suggest_source(label, label, document_type, text)
        regions.append(
            _region_values(
                path=f"{ref.path}#label:{match_index}",
                label=label,
                source=source,
                field_type=field_type,
                confidence=0.78 if gap else 0.7,
                region_data={
                    "kind": "label_value",
                    "start": match.end(),
                    "end": end,
                    "label_token": match.group(0),
                },
                ordinal=ordinal + len(regions),
                repeatable=document_type == DocumentType.PHOTO_DOCUMENT,
            )
        )
    return regions


def analyze_template(db: Session, version_id: uuid.UUID) -> None:
    version = db.get(TemplateVersion, version_id)
    if not version:
        raise ValueError("模板版本不存在")
    template = db.get(Template, version.template_id)
    if not template:
        raise ValueError("模板不存在")
    version.status = ProcessingStatus.RUNNING
    version.error_message = None
    db.commit()
    try:
        data = get_storage().read_bytes(version.normalized_object_key)
        document = Document(io.BytesIO(data))
        regions: list[dict[str, Any]] = []
        seen_paths: set[str] = set()
        empty_body_refs: list[ParagraphRef] = []
        ordinal = 0
        for ref in _paragraph_refs(document):
            text = ref.paragraph.text
            if ref.path.startswith("body:p:") and not text.strip():
                empty_body_refs.append(ref)
            for match_index, match in enumerate(PLACEHOLDER.finditer(text)):
                key = match.group(1)
                path = f"{ref.path}#placeholder:{match_index}"
                source, field_type = _suggest_source(key, key, template.document_type)
                values = _region_values(
                    path=path,
                    label=key,
                    source=source,
                    field_type=field_type,
                    confidence=1.0,
                    region_data={"token": match.group(0), "kind": "placeholder"},
                    ordinal=ordinal,
                )
                values["field_key"] = key
                regions.append(values)
                seen_paths.add(ref.path)
                ordinal += 1
            blank = BLANK.search(text)
            if blank and ref.path not in seen_paths:
                label = text[: blank.start()].strip(" ：:") or ref.context_label or "待填写字段"
                source, field_type = _suggest_source(label, label, template.document_type, text)
                regions.append(
                    _region_values(
                        path=f"{ref.path}#blank",
                        label=label,
                        source=source,
                        field_type=field_type,
                        confidence=0.8,
                        region_data={"token": blank.group(0), "kind": "blank"},
                        ordinal=ordinal,
                    )
                )
                seen_paths.add(ref.path)
                ordinal += 1
            photo_placeholder = PHOTO_PLACEHOLDER.search(text)
            if photo_placeholder and ref.path not in seen_paths:
                regions.append(
                    _region_values(
                        path=f"{ref.path}#photo",
                        label="照片",
                        source="document.photo_entries",
                        field_type="image",
                        confidence=0.98,
                        region_data={
                            "token": photo_placeholder.group(0),
                            "kind": "picture_placeholder",
                        },
                        ordinal=ordinal,
                        repeatable=True,
                    )
                )
                seen_paths.add(ref.path)
                ordinal += 1
            if ref.path not in seen_paths and text.strip():
                label_regions = _label_regions(ref, template.document_type, ordinal)
                if label_regions:
                    regions.extend(label_regions)
                    ordinal += len(label_regions)
                    seen_paths.add(ref.path)
            if not text.strip() and ref.context_label and ref.path not in seen_paths:
                label = ref.context_label.strip(" ：:")
                source, field_type = _suggest_source(
                    label, label, template.document_type, ref.context_label
                )
                regions.append(
                    _region_values(
                        path=f"{ref.path}#empty_cell",
                        label=label,
                        source=source,
                        field_type=field_type,
                        confidence=0.75,
                        region_data={"kind": "empty_cell"},
                        ordinal=ordinal,
                    )
                )
                seen_paths.add(ref.path)
                ordinal += 1
        if template.document_type == DocumentType.INTERVIEW_TRANSCRIPT and not any(
            region["source_path"] == "transcript.confirmed_text" for region in regions
        ):
            available = [ref for ref in empty_body_refs if ref.path not in seen_paths]
            if available:
                ref = available[-1]
                regions.append(
                    _region_values(
                        path=f"{ref.path}#interview_body",
                        label="询问问答正文",
                        source="transcript.confirmed_text",
                        field_type="multiline",
                        confidence=0.68,
                        region_data={"kind": "empty_paragraph"},
                        ordinal=ordinal,
                    )
                )
        existing = {
            region.path: region
            for region in db.scalars(
                select(TemplateRegion).where(TemplateRegion.template_version_id == version.id)
            )
        }
        for values in regions:
            region = existing.get(values["path"])
            if not region:
                db.add(TemplateRegion(template_version_id=version.id, **values))
                continue
            legacy_source = bool(
                region.source_path and re.fullmatch(r"custom\.[0-9a-f]{12}", region.source_path)
            )
            preserve_mapping = region.confirmed and not legacy_source
            region.region_data = values["region_data"]
            region.ordinal = values["ordinal"]
            region.repeatable = values["repeatable"]
            region.confidence = values["confidence"]
            if not preserve_mapping:
                region.label = values["label"]
                region.field_key = values["field_key"]
                region.source_path = values["source_path"]
                region.field_type = values["field_type"]
        detected_paths = {values["path"] for values in regions}
        for path, region in existing.items():
            if path in detected_paths:
                continue
            in_use = db.scalar(
                select(DocumentField.id).where(DocumentField.region_id == region.id).limit(1)
            )
            if not in_use:
                db.delete(region)
        version.status = ProcessingStatus.SUCCEEDED
    except Exception as exc:
        version.status = ProcessingStatus.FAILED
        version.error_message = str(exc)[:4000]
    db.commit()


def _resolve_paragraph(document: DocumentObject, path: str) -> Paragraph:
    base = path.split("#", 1)[0]
    parts = base.split(":")
    if parts[0] == "body" and parts[1] == "p":
        return document.paragraphs[int(parts[2])]
    if parts[0] == "body" and parts[1] == "table":
        return (
            document.tables[int(parts[2])]
            .rows[int(parts[4])]
            .cells[int(parts[6])]
            .paragraphs[int(parts[8])]
        )
    if parts[0] in {"header", "footer"}:
        section = document.sections[int(parts[1])]
        area = section.header if parts[0] == "header" else section.footer
        return area.paragraphs[int(parts[3])]
    raise ValueError(f"无法定位模板区域：{path}")


def _replace_across_runs(paragraph: Paragraph, token: str, replacement: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(replacement)
        return
    full = "".join(run.text for run in runs)
    start = full.find(token)
    if start < 0:
        raise ValueError(f"模板内容已变化，未找到填充标记：{token}")
    end = start + len(token)
    cursor = 0
    inserted = False
    for run in runs:
        run_start, run_end = cursor, cursor + len(run.text)
        if run_end <= start or run_start >= end:
            cursor = run_end
            continue
        left = run.text[: max(0, start - run_start)]
        right = run.text[max(0, end - run_start) :] if end < run_end else ""
        if not inserted:
            run.text = left + replacement + right
            inserted = True
        else:
            run.text = right
        cursor = run_end


def _replace_range(paragraph: Paragraph, start: int, end: int, replacement: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(replacement)
        return
    full = "".join(run.text for run in runs)
    if start < 0 or end < start or end > len(full):
        raise ValueError("模板填写区域的位置已经变化")

    def locate(offset: int) -> tuple[int, int]:
        cursor = 0
        for index, run in enumerate(runs):
            next_cursor = cursor + len(run.text)
            if offset <= next_cursor:
                return index, offset - cursor
            cursor = next_cursor
        return len(runs) - 1, len(runs[-1].text)

    start_index, start_offset = locate(start)
    end_index, end_offset = locate(end)
    if start_index == end_index:
        run = runs[start_index]
        run.text = run.text[:start_offset] + replacement + run.text[end_offset:]
        return
    runs[start_index].text = runs[start_index].text[:start_offset] + replacement
    for index in range(start_index + 1, end_index):
        runs[index].text = ""
    runs[end_index].text = runs[end_index].text[end_offset:]


def _repeat_path(
    path: str,
    block_index: int,
    paragraph_count: int,
    table_count: int,
) -> str:
    if block_index == 0 or not path.startswith("body:"):
        return path
    base, separator, suffix = path.partition("#")
    parts = base.split(":")
    if parts[1] == "p":
        parts[2] = str(int(parts[2]) + block_index * paragraph_count)
    elif parts[1] == "table":
        parts[2] = str(int(parts[2]) + block_index * table_count)
    repeated = ":".join(parts)
    return f"{repeated}{separator}{suffix}" if separator else repeated


def _repeat_document_body(document: DocumentObject, block_count: int) -> tuple[int, int]:
    paragraph_count = len(document.paragraphs)
    table_count = len(document.tables)
    if block_count <= 1:
        return paragraph_count, table_count
    body = document.element.body
    section_properties = body.sectPr
    block_elements = [
        deepcopy(element) for element in body.iterchildren() if element is not section_properties
    ]
    for _ in range(block_count - 1):
        for element in block_elements:
            section_properties.addprevious(deepcopy(element))
    return paragraph_count, table_count


def fill_template(
    template_data: bytes,
    fields: list[tuple[TemplateRegion, Any]],
    image_loader: callable | None = None,
) -> tuple[bytes, list[str]]:
    document = Document(io.BytesIO(template_data))
    warnings: list[str] = []
    photo_values = next(
        (
            raw_value
            for region, raw_value in fields
            if region.source_path == "document.photo_entries" and isinstance(raw_value, list)
        ),
        [],
    )
    block_count = max(1, len(photo_values))
    paragraph_count, table_count = _repeat_document_body(document, block_count)
    ordered_fields = sorted(fields, key=lambda item: item[0].ordinal, reverse=True)
    for block_index in range(block_count):
        for region, raw_value in ordered_fields:
            if block_index and not region.path.startswith("body:"):
                continue
            try:
                path = _repeat_path(
                    region.path,
                    block_index,
                    paragraph_count,
                    table_count,
                )
                paragraph = _resolve_paragraph(document, path)
                kind = region.region_data.get("kind")
                if region.source_path == "document.photo_entries" and photo_values:
                    current_value: Any = [photo_values[block_index]]
                elif isinstance(raw_value, list) and len(raw_value) == len(photo_values):
                    current_value = raw_value[block_index]
                else:
                    current_value = raw_value
                if region.field_type == "image":
                    token = region.region_data.get("token")
                    if token:
                        _replace_across_runs(paragraph, token, "")
                    values = current_value if isinstance(current_value, list) else [current_value]
                    if image_loader:
                        for item in values:
                            object_key = (
                                item.get("object_key")
                                if isinstance(item, dict)
                                else str(item or "")
                            )
                            if not object_key:
                                continue
                            with tempfile.NamedTemporaryFile(suffix=".jpg") as image_file:
                                image_file.write(image_loader(object_key))
                                image_file.flush()
                                paragraph.add_run().add_picture(image_file.name, width=Inches(2.3))
                    continue
                value = "" if current_value is None else str(current_value)
                token = region.region_data.get("token")
                if kind == "label_value":
                    _replace_range(
                        paragraph,
                        int(region.region_data["start"]),
                        int(region.region_data["end"]),
                        value,
                    )
                elif kind == "empty_paragraph":
                    _replace_range(paragraph, 0, len(paragraph.text), value)
                elif token:
                    _replace_across_runs(paragraph, token, value)
                elif paragraph.runs:
                    paragraph.runs[0].text = value
                else:
                    paragraph.add_run(value)
            except Exception as exc:
                suffix = f"（第 {block_index + 1} 张）" if block_count > 1 else ""
                warnings.append(f"{region.label}{suffix}：{exc}")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), warnings


def render_docx_preview(data: bytes) -> bytes | None:
    with tempfile.TemporaryDirectory() as temp_dir:
        source = Path(temp_dir) / "document.docx"
        source.write_bytes(data)
        try:
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", temp_dir, source],
                check=True,
                capture_output=True,
                timeout=90,
            )
        except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
            return None
        pdf = Path(temp_dir) / "document.pdf"
        return pdf.read_bytes() if pdf.exists() else None
