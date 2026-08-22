"""文书生成与渲染编排（PLAN 阶段 5）。

- 共享字段从检查单点获取；origin=user 的字段重新生成时保留（PRD §10）；
- 照片文书：逐帧字段（反映内容/说明）按关键帧， 同源字段只维护一次；
- 询问笔录：问/答归属优先取人工标注（segment.qa_role），否则 LLM 语义预标候选（澄清 §6）；
- 渲染：模板字段映射 → 6 类填充操作 → 下划线规则按文书类型 → 页脚域刷新。
"""
import re
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
from sqlalchemy.orm import Session
from ..core.config import MEDIA_DIR, EXPORT_DIR, settings
from ..core.models import (Document, DocumentField, Inspection, Template, TemplateField,
                           Fact, Keyframe, Transcript, TranscriptSegment)
from ..core import ai
from . import docx_fill

DOC_TYPE_NAME = {"record": "消防监督检查记录", "photo": "现场照片", "interrogation": "询问笔录"}
LOCAL_TZ = ZoneInfo("Asia/Shanghai")


def _local(dt: datetime) -> datetime:
    return dt.astimezone(LOCAL_TZ) if dt.tzinfo else dt.replace(tzinfo=ZoneInfo("UTC")).astimezone(LOCAL_TZ)
CHECK_TYPES = ["大型群众性活动举办前的检查", "建设工程施工现场检查", "对举报投诉的核查", "复查",
               "申请恢复施工、使用、生产、经营的检查", "申请解除临时查封的检查", "其它检查"]


def default_name(inspection: Inspection, doc_type: str) -> str:
    unit = inspection.unit_name or "（未填单位）"
    return f"{unit}{DOC_TYPE_NAME.get(doc_type, '文书')}"


def _get_field(db: Session, doc: Document, label: str) -> DocumentField | None:
    return next((f for f in doc.fields if f.field_label == label and not f.frame_key), None)


def _upsert(db: Session, doc: Document, label: str, value: str, per_frame=False, frame_key="",
            force=False):
    """写入字段值；origin=user 且非 force 时保留人工修改（PRD §10）。"""
    f = next((x for x in doc.fields if x.field_label == label and x.frame_key == frame_key), None)
    if f is None:
        f = DocumentField(document_id=doc.id, field_label=label, per_frame=per_frame, frame_key=frame_key)
        db.add(f)
        doc.fields.append(f)
    if f.origin == "user" and not force and f.value:
        return f
    f.value, f.origin = value, "ai"
    return f


def _confirmed_facts(db: Session, inspection_id: int) -> list[str]:
    return [f.content for f in db.query(Fact).filter_by(inspection_id=inspection_id, status="confirmed")]


RECORD_SYSTEM = """你是消防监督检查文书起草助手。根据已确认的检查事实，撰写"检查内容和情况"栏目文字。
要求：客观、简洁、公文体；按时间顺序描述检查过程与发现的问题；不得编造事实中没有的内容。
输出 JSON：{"content": "..."}"""


def generate_values(db: Session, doc: Document, frame_ids: list[str] | None = None):
    """按文书类型生成/刷新字段值（AI 初稿）。frame_ids：照片文书选帧集合（重建逐帧字段）。"""
    insp = db.get(Inspection, doc.inspection_id)
    facts = _confirmed_facts(db, insp.id)

    if doc.doc_type == "record":
        _upsert(db, doc, "单位（场所）名称", insp.unit_name)
        _upsert(db, doc, "法定代表/主要负责人", insp.legal_person)
        _upsert(db, doc, "地址", insp.address)
        _upsert(db, doc, "联系人及联系电话", insp.contact)
        if insp.check_type:
            _upsert(db, doc, "检查类型", insp.check_type)
        if insp.inspected_at:
            t = _local(insp.inspected_at)
            _upsert(db, doc, "检查时间", f"{t.year}|{t.month}|{t.day}")
        content = ""
        if facts and settings.ai_available("llm"):
            try:
                r = ai.llm_json(RECORD_SYSTEM, "已确认的检查事实：\n" + "\n".join(f"- {f}" for f in facts))
                content = r.get("content", "")
            except ai.AIError:
                content = "\n".join(facts)  # 降级：罗列事实，人工修改
        elif facts:
            content = "\n".join(facts)
        _upsert(db, doc, "检查内容和情况", content)

    elif doc.doc_type == "photo":
        _upsert(db, doc, "拍摄地点", insp.address)
        _upsert(db, doc, "摄制人", insp.checker_names)
        _upsert(db, doc, "确认人", insp.contact.split(" ")[0] if insp.contact else "")
        if insp.inspected_at:
            _upsert(db, doc, "拍摄时间", _local(insp.inspected_at).strftime("%Y年%m月%d日"))
        # 逐帧字段：选帧集合（入参优先，否则沿用已有）
        if frame_ids is not None:
            db.query(DocumentField).filter(DocumentField.document_id == doc.id,
                                           DocumentField.frame_key != "").delete(synchronize_session=False)
            db.flush()
            db.expire(doc, ["fields"])
            for fk in frame_ids:
                _upsert(db, doc, "反映内容", "", per_frame=True, frame_key=fk)
                _upsert(db, doc, "说明", "", per_frame=True, frame_key=fk)
        frame_keys = sorted({f.frame_key for f in doc.fields if f.frame_key})
        for fk in frame_keys:
            kf = db.get(Keyframe, int(fk)) if fk.isdigit() else None
            caption = kf.caption if kf else ""
            _upsert(db, doc, "反映内容", caption, per_frame=True, frame_key=fk)
            _upsert(db, doc, "说明", _one_sentence(caption), per_frame=True, frame_key=fk)

    elif doc.doc_type == "interrogation":
        if insp.inspected_at:
            _upsert(db, doc, "时间（起止）", _local(insp.inspected_at).strftime("%Y年%m月%d日 %H:%M") + "|")
        _upsert(db, doc, "地点", insp.address)
        body = _build_qa_body(db, insp.id)
        if body:
            _upsert(db, doc, "问答正文", body, force=True)  # 正文来源变化大，始终刷新（人工版在编辑页另存）
    db.flush()


def _one_sentence(text: str, limit: int = 50) -> str:
    text = re.split(r"[。！？\n]", (text or "").strip())[0]
    return text[:limit]


QA_SYSTEM = """你是询问笔录整理助手。把消防检查现场对话转写整理成"问/答"轮次。
要求：
1. 检查人员的发问标"问"，被检查单位人员的回答/陈述标"答"；按语义（疑问句/陈述句）判断，不确定的归"答"。
2. 修正明显口误与 ASR 错字（如"24h"→"24小时"），但不得添加原对话没有的内容。
3. 合并零碎语句为通顺的书面句，保持原意。
输出 JSON：{"rounds": [{"role": "问"|"答", "text": "..."}]}"""


def _build_qa_body(db: Session, inspection_id: int) -> str:
    """优先人工标注的窗口分段；否则 LLM 语义归属整段转写（候选，PRD §7.3 不得猜测身份）。"""
    # 人工标注优先：本检查所有窗口分段中带 qa_role 的
    from ..core.models import Material
    segs = (db.query(TranscriptSegment).join(Transcript)
            .join(Material, Transcript.material_id == Material.id)
            .filter(Transcript.kind == "window", TranscriptSegment.qa_role.isnot(None),
                    Material.inspection_id == inspection_id).all())
    if segs:
        lines = []
        for s in sorted(segs, key=lambda x: (x.material_id, x.seq)):
            if s.text.strip():
                lines.append(f"{s.qa_role}：{s.text.strip()}")
        if lines:
            return "\n".join(lines)
    # LLM 候选
    fulls = (db.query(Transcript).join(Transcript.material)
             .filter_by(inspection_id=inspection_id).filter(Transcript.kind == "full").all())
    texts = [t.segments[0].text for t in fulls if t.segments and t.segments[0].text.strip()]
    if not texts or not settings.ai_available("llm"):
        return "\n".join(texts)
    try:
        r = ai.llm_json(QA_SYSTEM, "现场对话转写：\n" + "\n\n".join(texts))
        return "\n".join(f"{x.get('role','答')}：{x.get('text','').strip()}"
                         for x in r.get("rounds", []) if x.get("text"))
    except ai.AIError:
        return "\n".join(texts)  # 降级：原始转写，人工整理


def render_document(db: Session, doc: Document) -> str:
    """按模板映射渲染 docx，返回产物路径（相对 data_dir）。"""
    tpl = db.get(Template, doc.template_id)
    if not tpl:
        raise ValueError("文书未绑定模板")
    underline = doc.doc_type in ("photo", "interrogation")  # 记录表普通文本（TEMPLATE-ANALYSIS §2）
    fields: list[TemplateField] = [f for f in tpl.fields if f.enabled and f.confirmed]

    if doc.doc_type == "photo":
        return _render_photo(db, doc, tpl, fields, underline)

    ops = []
    for tf in fields:
        df = _get_field(db, doc, tf.field_label)
        value = (df.value if df else "").strip()
        if tf.fill_op == "generate_body":
            if value:
                ops.append({"fill_op": "generate_body", "anchor": tf.anchor, "value": value,
                            "underline": True})
            continue
        if not value:
            continue  # 留空保持模板空白（如手写签名）
        op = {"fill_op": tf.fill_op, "anchor": tf.anchor, "value": value, "underline": underline}
        if tf.fill_op == "checkbox":
            op["choices"] = tf.options.get("choices", [])
        ops.append(op)
    out = EXPORT_DIR / f"doc_{doc.id}.docx"
    docx_fill.fill_template(str(MEDIA_DIR.parent / tpl.file_path), str(out), ops)
    doc.file_path = str(out.relative_to(MEDIA_DIR.parent))
    db.flush()
    return doc.file_path


def _render_photo(db: Session, doc: Document, tpl: Template, fields: list[TemplateField], underline: bool) -> str:
    frame_keys = sorted({f.frame_key for f in doc.fields if f.frame_key})
    if not frame_keys:
        raise ValueError("照片文书尚未选择照片/关键帧")
    p_indexes = [f.anchor["p_index"] for f in fields if f.anchor.get("kind") == "paragraph"]
    start = min(p_indexes)
    import docx as _dx
    raw = _dx.Document(str(MEDIA_DIR.parent / tpl.file_path))
    end = max(p_indexes) + 1
    while end < len(raw.paragraphs) and not raw.paragraphs[end].text.strip():
        end += 1  # 含分隔空段（P8）
    block_len = end - start

    def shifted(anchor, k):
        a = dict(anchor)
        if a.get("kind") == "paragraph":
            a["p_index"] += k * block_len
        return a

    ops = []
    for k, fk in enumerate(frame_keys):
        kf = db.get(Keyframe, int(fk)) if str(fk).isdigit() else None
        for tf in fields:
            if tf.fill_op == "replace_placeholder_image":
                if kf:
                    ops.append({"fill_op": "replace_placeholder_image", "anchor": shifted(tf.anchor, k),
                                "value": str(MEDIA_DIR.parent / kf.image_path)})
                continue
            if tf.per_frame:
                df = next((x for x in doc.fields if x.field_label == tf.field_label and x.frame_key == fk), None)
            else:
                df = _get_field(db, doc, tf.field_label)
            value = (df.value if df else "").strip()
            if value:
                ops.append({"fill_op": tf.fill_op, "anchor": shifted(tf.anchor, k),
                            "value": value, "underline": underline})
    out = EXPORT_DIR / f"doc_{doc.id}.docx"
    docx_fill.fill_template(str(MEDIA_DIR.parent / tpl.file_path), str(out), ops,
                            photo_block={"start": start, "len": block_len, "n_frames": len(frame_keys)})
    doc.file_path = str(out.relative_to(MEDIA_DIR.parent))
    db.flush()
    return doc.file_path
