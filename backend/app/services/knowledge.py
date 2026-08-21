import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core.storage import get_storage
from app.models import Citation, KnowledgeChunk, KnowledgeDocument, LegalAnalysis
from app.services.ai import ai_provider, cosine_similarity


def extract_document(data: bytes, filename: str) -> list[tuple[str, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = data.decode("utf-8-sig")
        return [("全文", text)] if text.strip() else []
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return [
            (f"第 {index} 页", text)
            for index, page in enumerate(reader.pages, 1)
            if (text := (page.extract_text() or "").strip())
        ]
    if suffix == ".docx":
        document = Document(io.BytesIO(data))
        sections: list[tuple[str, str]] = []
        current_title = "正文"
        buffer: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                if buffer:
                    sections.append((current_title, "\n".join(buffer)))
                    buffer = []
                current_title = text
            else:
                buffer.append(text)
        for table in document.tables:
            for row in table.rows:
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    buffer.append(text)
        if buffer:
            sections.append((current_title, "\n".join(buffer)))
        return sections
    raise ValueError("仅支持 PDF、DOCX 和 TXT")


def split_chunks(sections: list[tuple[str, str]], max_chars: int = 1000) -> list[tuple[str, str]]:
    chunks: list[tuple[str, str]] = []
    for locator, text in sections:
        paragraphs = [item.strip() for item in re.split(r"\n+", text) if item.strip()]
        buffer = ""
        for paragraph in paragraphs:
            if buffer and len(buffer) + len(paragraph) + 1 > max_chars:
                chunks.append((locator, buffer))
                buffer = ""
            if len(paragraph) > max_chars:
                for start in range(0, len(paragraph), max_chars):
                    piece = paragraph[start : start + max_chars]
                    if piece:
                        chunks.append((locator, piece))
            else:
                buffer = f"{buffer}\n{paragraph}".strip()
        if buffer:
            chunks.append((locator, buffer))
    return chunks


def index_document(db: Session, document_id: str) -> None:
    document = db.get(KnowledgeDocument, document_id)
    if not document:
        raise ValueError("知识文档不存在")
    try:
        sections = extract_document(
            get_storage().read_bytes(document.storage_key), document.filename
        )
        chunks = split_chunks(sections)
        if not chunks:
            raise ValueError("未提取到可检索文本；扫描版文档暂不支持，请先进行 OCR")
        vectors = ai_provider.embeddings([content for _, content in chunks])
        db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.document_id == document.id))
        for ordinal, ((locator, content), vector) in enumerate(zip(chunks, vectors, strict=True)):
            db.add(
                KnowledgeChunk(
                    document_id=document.id,
                    ordinal=ordinal,
                    locator=locator,
                    content=content,
                    embedding=vector,
                )
            )
        document.indexing_status = "ready"
        document.indexing_error = None
        db.commit()
    except Exception as exc:
        document.indexing_status = "failed"
        document.indexing_error = str(exc)
        db.commit()
        raise


def _lexical_score(question: str, content: str) -> float:
    lowered_question = question.lower()
    chinese = "".join(re.findall(r"[\u4e00-\u9fff]", lowered_question))
    terms = {chinese[index : index + 2] for index in range(max(0, len(chinese) - 1))}
    terms.update(re.findall(r"[a-z0-9_]{2,}", lowered_question))
    if not terms:
        return 0.0
    lowered = content.lower()
    return sum(1 for term in terms if term in lowered) / len(terms)


def query_knowledge(
    db: Session,
    *,
    question: str,
    user_id: str,
    inspection_id: str | None = None,
    fact_id: str | None = None,
    limit: int = 5,
) -> tuple[LegalAnalysis, list[Citation]]:
    base_query = (
        select(KnowledgeChunk, KnowledgeDocument)
        .join(KnowledgeDocument, KnowledgeDocument.id == KnowledgeChunk.document_id)
        .where(KnowledgeDocument.enabled.is_(True), KnowledgeDocument.indexing_status == "ready")
    )
    query_vector = ai_provider.embeddings([question])[0]
    if db.bind.dialect.name == "postgresql":
        vector_score = 1 - KnowledgeChunk.embedding.cosine_distance(query_vector)
        text_score = func.similarity(KnowledgeChunk.content, question)
        hybrid_score = 0.55 * vector_score + 0.45 * text_score
        ranked = db.execute(base_query.order_by(hybrid_score.desc()).limit(limit)).all()
    else:
        rows = db.execute(base_query).all()
        ranked = sorted(
            rows,
            key=lambda row: (
                0.55 * cosine_similarity(query_vector, list(row[0].embedding))
                + 0.45 * _lexical_score(question, row[0].content)
            ),
            reverse=True,
        )[:limit]

    if not ranked or all(_lexical_score(question, chunk.content) == 0 for chunk, _ in ranked):
        characterization = "未找到足够依据"
        reasoning = "当前知识库中没有检索到能够可靠支持判断的原文，请补充材料或调整描述。"
        selected = []
    else:
        selected = ranked
        context = "\n\n".join(
            f"[CHUNK:{chunk.id}] {doc.title} {chunk.locator or ''}\n{chunk.content}"
            for chunk, doc in selected
        )
        generated = ai_provider.chat_json(
            system=(
                "你是消防法规检索助手。只能依据提供的 CHUNK 作答，不作最终执法判断。"
                "返回 JSON：characterization、reasoning、chunk_ids。chunk_ids 只能使用输入中的 ID。"
            ),
            user=f"问题：{question}\n\n检索材料：\n{context}",
        )
        valid_ids = {chunk.id for chunk, _ in selected}
        cited_ids = [item for item in (generated or {}).get("chunk_ids", []) if item in valid_ids]
        if generated and cited_ids:
            characterization = str(generated.get("characterization") or "可能涉及消防安全问题")
            reasoning = str(generated.get("reasoning") or "请结合引用原文人工判断。")
            selected = [row for row in selected if row[0].id in cited_ids]
        else:
            characterization = "可能涉及与现场描述相关的消防安全要求"
            reasoning = "以下为知识库中检索到的相关原文，具体适用性须由执法人员核对。"

    analysis = LegalAnalysis(
        inspection_id=inspection_id,
        fact_id=fact_id,
        question=question,
        characterization=characterization,
        reasoning=reasoning,
        created_by=user_id,
    )
    db.add(analysis)
    db.flush()
    citations = []
    for chunk, document in selected:
        citation = Citation(
            analysis_id=analysis.id,
            document_id=document.id,
            chunk_id=chunk.id,
            excerpt=chunk.content,
            locator=chunk.locator,
        )
        db.add(citation)
        citations.append(citation)
    db.commit()
    return analysis, citations
