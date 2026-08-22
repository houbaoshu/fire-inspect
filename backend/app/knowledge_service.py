from __future__ import annotations

import io
import math
import re
import subprocess
import tempfile
import uuid
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from .ai import get_ai
from .models import (
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeDocumentVersion,
    ProcessingStatus,
    ProcessingStep,
    StepKind,
)
from .storage import get_storage


def _paragraphs_from_file(
    data: bytes, content_type: str, filename: str
) -> list[tuple[int | None, str]]:
    lower = filename.lower()
    if lower.endswith(".pdf") or content_type == "application/pdf":
        reader = PdfReader(io.BytesIO(data))
        paragraphs: list[tuple[int | None, str]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            for paragraph in re.split(r"\n\s*\n|\n(?=[\u4e00-\u9fffA-Z0-9])", text):
                if paragraph.strip():
                    paragraphs.append((page_number, paragraph.strip()))
        return paragraphs
    if lower.endswith(".docx"):
        document = Document(io.BytesIO(data))
        paragraphs = [(None, p.text.strip()) for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    paragraphs.append((None, text))
        return paragraphs
    if lower.endswith(".txt") or content_type.startswith("text/"):
        text = data.decode("utf-8", errors="replace")
        return [(None, p.strip()) for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    raise ValueError("知识文档仅支持 PDF、DOCX 和 TXT")


def _chunk(
    paragraphs: list[tuple[int | None, str]], limit: int = 900
) -> list[tuple[int | None, str]]:
    result: list[tuple[int | None, str]] = []
    current: list[str] = []
    current_page: int | None = None
    size = 0
    for page, paragraph in paragraphs:
        if current and (size + len(paragraph) > limit or page != current_page):
            result.append((current_page, "\n".join(current)))
            current, size = [], 0
        current_page = page
        current.append(paragraph)
        size += len(paragraph)
    if current:
        result.append((current_page, "\n".join(current)))
    return result


def _ocr_pdf(data: bytes) -> list[tuple[int | None, str]]:
    paragraphs: list[tuple[int | None, str]] = []
    with tempfile.TemporaryDirectory() as temp_dir:
        source = Path(temp_dir) / "source.pdf"
        source.write_bytes(data)
        prefix = Path(temp_dir) / "page"
        subprocess.run(
            ["pdftoppm", "-jpeg", "-r", "150", str(source), str(prefix)],
            check=True,
            capture_output=True,
        )
        for page_number, image in enumerate(sorted(Path(temp_dir).glob("page-*.jpg")), start=1):
            text = get_ai().extract(image.read_bytes()).strip()
            if text:
                paragraphs.append((page_number, text))
    return paragraphs


def parse_knowledge_version(db: Session, version_id: uuid.UUID) -> None:
    version = db.get(KnowledgeDocumentVersion, version_id)
    if not version:
        raise ValueError("知识文档版本不存在")
    is_pdf = version.filename.lower().endswith(".pdf") or version.content_type == "application/pdf"
    step_kinds = [StepKind.PARSE, StepKind.OCR_VISION, StepKind.CHUNKING, StepKind.EMBEDDING]
    steps = {
        step.kind: step
        for step in db.scalars(
            select(ProcessingStep).where(
                ProcessingStep.subject_type == "knowledge_version",
                ProcessingStep.subject_id == version.id,
            )
        )
    }
    for kind in step_kinds:
        if kind not in steps:
            steps[kind] = ProcessingStep(
                subject_type="knowledge_version",
                subject_id=version.id,
                kind=kind,
                algorithm_version="v1",
            )
            db.add(steps[kind])
    version.status = ProcessingStatus.RUNNING
    version.error_message = None
    db.commit()
    try:
        parse_step = steps[StepKind.PARSE]
        parse_step.status = ProcessingStatus.RUNNING
        parse_step.attempt_count += 1
        db.commit()
        data = get_storage().read_bytes(version.object_key)
        paragraphs = _paragraphs_from_file(data, version.content_type, version.filename)
        parse_step.status = ProcessingStatus.SUCCEEDED
        parse_step.artifact = {"paragraph_count": len(paragraphs)}
        ocr_step = steps[StepKind.OCR_VISION]
        if not paragraphs:
            if not is_pdf:
                raise ValueError("未解析出文字")
            ocr_step.status = ProcessingStatus.RUNNING
            ocr_step.attempt_count += 1
            db.commit()
            paragraphs = _ocr_pdf(data)
            if not paragraphs:
                raise ValueError("扫描版 PDF 未识别出文字，请配置 OCR Provider 后重试")
            ocr_step.status = ProcessingStatus.SUCCEEDED
            ocr_step.artifact = {"paragraph_count": len(paragraphs)}
        else:
            ocr_step.status = ProcessingStatus.SUCCEEDED
            ocr_step.artifact = {"skipped": True, "reason": "document_has_text_layer"}
        chunk_step = steps[StepKind.CHUNKING]
        chunk_step.status = ProcessingStatus.RUNNING
        chunk_step.attempt_count += 1
        db.commit()
        chunks = _chunk(paragraphs)
        chunk_step.status = ProcessingStatus.SUCCEEDED
        chunk_step.artifact = {"chunk_count": len(chunks)}
        embedding_step = steps[StepKind.EMBEDDING]
        embedding_step.status = ProcessingStatus.RUNNING
        embedding_step.attempt_count += 1
        db.commit()
        embeddings = get_ai().embed([text for _, text in chunks])
        db.execute(delete(KnowledgeChunk).where(KnowledgeChunk.version_id == version.id))
        for index, ((page, text), embedding) in enumerate(zip(chunks, embeddings, strict=True)):
            db.add(
                KnowledgeChunk(
                    version_id=version.id,
                    anchor=f"chunk-{index + 1}",
                    page_number=page,
                    text=text,
                    embedding=embedding,
                )
            )
        embedding_step.status = ProcessingStatus.SUCCEEDED
        embedding_step.artifact = {"vector_count": len(embeddings), "dimensions": 32}
        version.status = ProcessingStatus.SUCCEEDED
    except Exception as exc:
        running = next(
            (step for step in steps.values() if step.status == ProcessingStatus.RUNNING), None
        )
        if running:
            running.status = ProcessingStatus.FAILED
            running.error_code = type(exc).__name__
            running.error_message = str(exc)[:4000]
        version.status = ProcessingStatus.FAILED
        version.error_message = str(exc)[:4000]
    db.commit()


def _tokens(text: str) -> list[str]:
    lowered = text.lower()
    words = re.findall(r"[a-z0-9_]+", lowered)
    chinese = "".join(re.findall(r"[\u4e00-\u9fff]", lowered))
    return words + [chinese[index : index + 2] for index in range(max(0, len(chinese) - 1))]


def _lexical_score(query: str, text: str) -> float:
    query_tokens = Counter(_tokens(query))
    text_tokens = Counter(_tokens(text))
    if not query_tokens:
        return 0.0
    overlap = sum(min(count, text_tokens[token]) for token, count in query_tokens.items())
    return overlap / sum(query_tokens.values())


def _cosine(left: list[float], right: list[float]) -> float:
    numerator = sum(a * b for a, b in zip(left, right, strict=False))
    denominator = math.sqrt(sum(a * a for a in left)) * math.sqrt(sum(b * b for b in right))
    return numerator / denominator if denominator else 0.0


def search_knowledge(db: Session, query: str, limit: int) -> dict[str, Any]:
    query_embedding = get_ai().embed([query])[0]
    rows = db.execute(
        select(KnowledgeChunk, KnowledgeDocumentVersion, KnowledgeDocument)
        .join(KnowledgeDocumentVersion, KnowledgeDocumentVersion.id == KnowledgeChunk.version_id)
        .join(KnowledgeDocument, KnowledgeDocument.id == KnowledgeDocumentVersion.document_id)
        .where(
            KnowledgeDocument.status == "active",
            KnowledgeDocumentVersion.status == ProcessingStatus.SUCCEEDED,
        )
    ).all()
    ranked: list[tuple[float, KnowledgeChunk, KnowledgeDocumentVersion, KnowledgeDocument]] = []
    for chunk, version, document in rows:
        vector = list(chunk.embedding or [])
        score = 0.55 * _lexical_score(query, chunk.text) + 0.45 * _cosine(query_embedding, vector)
        if score > 0.05:
            ranked.append((score, chunk, version, document))
    ranked.sort(key=lambda item: item[0], reverse=True)
    selected = ranked[:limit]
    sources = [{"text": chunk.text, "score": score} for score, chunk, _, _ in selected]
    answer = get_ai().answer_with_sources(query, sources)
    citations = [
        {
            "document_id": document.id,
            "version_id": version.id,
            "title": document.title,
            "category": document.category,
            "anchor": chunk.anchor,
            "page_number": chunk.page_number,
            "quote": chunk.text,
        }
        for _, chunk, version, document in selected
    ]
    return {**answer, "citations": citations}
